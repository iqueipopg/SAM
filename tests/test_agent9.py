# Copyright 2026 Ignacio Queipo de Llano Pérez-Gascón and the SAM authors
# SPDX-License-Identifier: Apache-2.0
"""Comprehensive tests for Agent9Explainer (TEST1).

A9 is a CPU-only post-processing agent that generates:
  - papers_raw.xlsx: flat Excel with one row per paper + pipeline decisions
  - explainability.md: Spanish-language traceability report

These tests cover:
  1. Excel schema (exact column count + headers).
  2. explainability.md: sections present, Spanish content, A2 sources (BUG4).
  3. Edge cases: 0 papers, None/empty fields, Unicode, malformed DOIs.
  4. Determinism: same input → same files (datetime mocked).
  5. Output dir safety: creates dir if missing, overwrites existing files.

Runs entirely offline: no GPU, no LLM, no network. Uses openpyxl.load_workbook
to introspect the generated Excel file.
"""

import copy
import hashlib
import os
import sys
from unittest.mock import patch

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))


# =========================================================================== #
#  HELPERS                                                                     #
# =========================================================================== #

def _load_workbook(path):
    from openpyxl import load_workbook
    return load_workbook(path)


def _read_md(path):
    with open(path, encoding="utf-8") as f:
        return f.read()


def _run_a9(pipeline_data, tmp_path, run_id="test_run_0001", subdir="a9_out"):
    """Instantiate A9 and run it against a pipeline dict."""
    from agents.agent9_explainer import Agent9Explainer

    out_dir = tmp_path / subdir
    agent = Agent9Explainer()
    return agent.run(pipeline_data, run_id=run_id, output_dir=str(out_dir))


# =========================================================================== #
#  1. EXCEL SCHEMA                                                             #
# =========================================================================== #

def test_a9_excel_has_expected_column_count(full_pipeline_data, a9_excel_columns, tmp_path):
    """papers_raw.xlsx must have exactly the number of columns declared in the
    agent source (27 at this iteration)."""
    res = _run_a9(full_pipeline_data, tmp_path)
    wb = _load_workbook(res["excel_path"])
    ws = wb.active
    # Row 1 is headers. Count non-empty cells.
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    headers = [h for h in headers if h is not None]

    assert len(headers) == len(a9_excel_columns), (
        f"Expected {len(a9_excel_columns)} columns, found {len(headers)}"
    )
    # 27 is the current contract as per agent9_explainer.py
    assert len(headers) == 27


def test_a9_excel_header_names_match(full_pipeline_data, a9_excel_columns, tmp_path):
    """Header names must match EXACTLY the fixture column list."""
    res = _run_a9(full_pipeline_data, tmp_path)
    wb = _load_workbook(res["excel_path"])
    ws = wb.active
    headers = [ws.cell(row=1, column=c).value for c in range(1, len(a9_excel_columns) + 1)]
    assert headers == a9_excel_columns


def test_a9_excel_row_count_matches_pipeline(full_pipeline_data, tmp_path):
    """Excel must contain one data row per unique PMID seen by the pipeline.
    In the fixture: 4 papers entered A3 (3 included + 1 excluded)."""
    res = _run_a9(full_pipeline_data, tmp_path)
    wb = _load_workbook(res["excel_path"])
    ws = wb.active
    # max_row counts header + data rows
    assert ws.max_row == 1 + 4  # header + 4 papers


def test_a9_excel_contains_pipeline_decisions(full_pipeline_data, tmp_path):
    """Spot-check: the Excel body should carry A3 decisions, A6 tools, and the
    exclusion stage for the excluded paper."""
    res = _run_a9(full_pipeline_data, tmp_path)
    wb = _load_workbook(res["excel_path"])
    ws = wb.active

    # Collect all values in the A3 Decision column (col 17)
    a3_decisions = [ws.cell(row=r, column=17).value for r in range(2, ws.max_row + 1)]
    assert "INCLUDE" in a3_decisions
    assert "EXCLUDE" in a3_decisions

    # A6 Tool column (col 23) - at least one RoB2
    a6_tools = [ws.cell(row=r, column=23).value for r in range(2, ws.max_row + 1)]
    assert any(t == "RoB2" for t in a6_tools)

    # Exclusion Stage column (col 27) - one paper excluded at title_abstract
    exclusion_stages = [ws.cell(row=r, column=27).value for r in range(2, ws.max_row + 1)]
    assert "title_abstract" in exclusion_stages


# =========================================================================== #
#  2. EXPLAINABILITY.MD (SPANISH REPORT)                                       #
# =========================================================================== #

def test_a9_report_has_all_six_sections(full_pipeline_data, tmp_path):
    """explainability.md must contain the 6 documented sections."""
    res = _run_a9(full_pipeline_data, tmp_path)
    md = _read_md(res["report_path"])

    expected_sections = [
        "## 1. Resumen General",
        "## 2. Trazabilidad por Etapa",
        "## 3. Razones de Exclusion por Paper",
        "## 4. Trazabilidad RAG",
        "## 5. Decisiones Borderline",
        "## 6. Procedencia de Datos por Paper Incluido",
    ]
    for sec in expected_sections:
        assert sec in md, f"Missing section: {sec}"


def test_a9_report_is_in_spanish(full_pipeline_data, tmp_path):
    """The report must contain typical Spanish keywords (regla #12 del proyecto)."""
    res = _run_a9(full_pipeline_data, tmp_path)
    md = _read_md(res["report_path"]).lower()

    # At least these Spanish-specific markers must appear
    spanish_markers = [
        "trazabilidad",
        "busqueda",
        "cribado",
        "calidad",
        "exclusion",
        "procedencia",
        "bases de datos",
    ]
    for marker in spanish_markers:
        assert marker in md, f"Spanish keyword missing: {marker}"


def test_a9_report_includes_a2_sources_when_provided(full_pipeline_data, tmp_path):
    """BUG4: when `a2_sources` is injected, A9 must list each database."""
    full_pipeline_data["a2_sources"] = [
        "pubmed", "europepmc", "semantic_scholar", "google_scholar"
    ]
    res = _run_a9(full_pipeline_data, tmp_path)
    md = _read_md(res["report_path"])

    assert "Bases de datos consultadas por A2" in md
    for db in ("pubmed", "europepmc", "semantic_scholar", "google_scholar"):
        assert db in md


def test_a9_report_contains_run_id(full_pipeline_data, tmp_path):
    """The report header must include the run_id passed in."""
    res = _run_a9(full_pipeline_data, tmp_path, run_id="run_unit_test_XYZ")
    md = _read_md(res["report_path"])
    assert "run_unit_test_XYZ" in md


# =========================================================================== #
#  3. EDGE CASES                                                               #
# =========================================================================== #

def test_a9_zero_included_papers(full_pipeline_data, tmp_path):
    """Pipeline with no included papers must not crash and must render a
    sensible message in section 6."""
    data = copy.deepcopy(full_pipeline_data)
    data["fulltext_retrieval"]["include"] = []
    data["fulltext_retrieval"]["exclude"] = []
    data["fulltext_retrieval"]["papers"] = []
    # Also clear screening.include so no paper is "included"
    data["screening"]["include"] = []
    data["screening"]["exclude"] = []
    data["screening"]["uncertain"] = []

    res = _run_a9(data, tmp_path, subdir="zero_out")

    # Excel generated with only headers
    wb = _load_workbook(res["excel_path"])
    ws = wb.active
    assert ws.max_row == 1, "Only header row should remain"

    # Report says "No hay papers incluidos." in section 6
    md = _read_md(res["report_path"])
    assert "No hay papers incluidos" in md


def test_a9_handles_none_fields_gracefully(full_pipeline_data, tmp_path):
    """Papers with None/empty fields must not raise: empty string or
    fallback should be written instead."""
    data = copy.deepcopy(full_pipeline_data)
    # Blank out several fields of the first screening paper
    paper = data["screening"]["include"][0]
    paper["title"] = None
    paper["authors"] = None
    paper["abstract"] = None

    # Also strip fulltext retrieval authors for first include
    ft0 = data["fulltext_retrieval"]["include"][0]
    ft0["authors"] = None
    ft0["journal"] = None

    res = _run_a9(data, tmp_path, subdir="none_out")

    # Excel should still open
    wb = _load_workbook(res["excel_path"])
    ws = wb.active
    assert ws.max_row >= 2

    # Report still valid
    md = _read_md(res["report_path"])
    assert "# Informe de Explicabilidad" in md


def test_a9_handles_unicode_in_title_and_authors(full_pipeline_data, tmp_path):
    """Non-ASCII / Unicode chars must propagate to both Excel and MD."""
    data = copy.deepcopy(full_pipeline_data)
    # Inject Unicode in screening paper 0 + matching fulltext entry
    unicode_title = "El tratamiento cognitivo en la depresión: estudio de García-López"
    unicode_authors = ["García-López R", "Müller K", "Tanaka Yōko"]

    data["screening"]["include"][0]["title"] = unicode_title
    data["screening"]["include"][0]["authors"] = unicode_authors
    data["fulltext_retrieval"]["include"][0]["title"] = unicode_title
    data["fulltext_retrieval"]["include"][0]["authors"] = unicode_authors

    res = _run_a9(data, tmp_path, subdir="unicode_out")

    # Excel contains the Unicode title
    wb = _load_workbook(res["excel_path"])
    ws = wb.active
    titles = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
    assert any(t == unicode_title for t in titles if t is not None)

    # MD encoded UTF-8 preserves accents
    md = _read_md(res["report_path"])
    assert "García-López" in md or "depresión" in md or "Müller" in md


def test_a9_handles_malformed_dois(full_pipeline_data, tmp_path):
    """Invalid / weird DOI strings must not cause an exception."""
    data = copy.deepcopy(full_pipeline_data)
    weird_dois = [
        "not-a-doi",
        "10.1000/",
        "",
        None,
        "https://doi.org/10.1000/valid",
    ]
    for i, ft_paper in enumerate(data["fulltext_retrieval"]["papers"]):
        ft_paper["doi"] = weird_dois[i % len(weird_dois)]

    res = _run_a9(data, tmp_path, subdir="doi_out")
    # If no exception, we're good; also verify files exist
    assert os.path.exists(res["excel_path"])
    assert os.path.exists(res["report_path"])


def test_a9_handles_missing_screening_stats(full_pipeline_data, tmp_path):
    """Missing `screening.stats` block must not raise; numbers fall back to 0."""
    data = copy.deepcopy(full_pipeline_data)
    data["screening"].pop("stats", None)
    res = _run_a9(data, tmp_path, subdir="nostats_out")
    md = _read_md(res["report_path"])
    # Should still contain the A3 section
    assert "### A3" in md


# =========================================================================== #
#  4. DETERMINISM                                                              #
# =========================================================================== #

def test_a9_is_deterministic_with_frozen_time(full_pipeline_data, tmp_path):
    """Same input + mocked datetime → same bytes for markdown.

    The Excel may include non-deterministic metadata (created-at timestamp
    embedded by openpyxl), so we compare only the markdown file.
    """
    from datetime import datetime as real_datetime

    class FrozenDatetime(real_datetime):
        @classmethod
        def now(cls, tz=None):
            return real_datetime(2026, 4, 21, 12, 0, 0)

    data1 = copy.deepcopy(full_pipeline_data)
    data2 = copy.deepcopy(full_pipeline_data)

    with patch("agents.agent9_explainer.datetime", FrozenDatetime):
        res1 = _run_a9(data1, tmp_path, subdir="det_out_1")

    with patch("agents.agent9_explainer.datetime", FrozenDatetime):
        res2 = _run_a9(data2, tmp_path, subdir="det_out_2")

    h1 = hashlib.sha256(open(res1["report_path"], "rb").read()).hexdigest()
    h2 = hashlib.sha256(open(res2["report_path"], "rb").read()).hexdigest()
    assert h1 == h2, "explainability.md should be byte-identical for same input"


# =========================================================================== #
#  5. OUTPUT DIR SAFETY                                                        #
# =========================================================================== #

def test_a9_creates_output_dir_when_missing(full_pipeline_data, tmp_path):
    """If output_dir does not exist, A9 must create it (exist_ok=True)."""
    from agents.agent9_explainer import Agent9Explainer

    missing = tmp_path / "does" / "not" / "exist" / "yet"
    assert not missing.exists()

    agent = Agent9Explainer()
    res = agent.run(full_pipeline_data, run_id="test_mkdir", output_dir=str(missing))

    assert missing.exists()
    assert os.path.exists(res["excel_path"])
    assert os.path.exists(res["report_path"])


def test_a9_overwrites_existing_files(full_pipeline_data, tmp_path):
    """Running A9 twice into the same dir must overwrite previous files
    (not append nor error out)."""
    from agents.agent9_explainer import Agent9Explainer

    out_dir = tmp_path / "overwrite_out"
    out_dir.mkdir()

    # Pre-seed with garbage content
    (out_dir / "papers_raw.xlsx").write_bytes(b"not a real xlsx")
    (out_dir / "explainability.md").write_text("stale content", encoding="utf-8")

    agent = Agent9Explainer()
    res = agent.run(full_pipeline_data, run_id="r1", output_dir=str(out_dir))

    # The markdown should no longer contain the stale content
    md = _read_md(res["report_path"])
    assert "stale content" not in md
    # The xlsx should now be a valid workbook
    wb = _load_workbook(res["excel_path"])
    assert wb.active.max_row >= 1


def test_a9_returns_expected_result_dict(full_pipeline_data, tmp_path):
    """A9.run() must return dict with `excel_path`, `report_path`, `stats`."""
    res = _run_a9(full_pipeline_data, tmp_path, subdir="result_dict_out")
    assert set(res.keys()) == {"excel_path", "report_path", "stats"}
    stats = res["stats"]
    for key in ("total_papers", "included", "excluded", "elapsed_s"):
        assert key in stats
    assert stats["total_papers"] == 4
    assert stats["included"] == 2  # fulltext.include has 2 papers
    assert stats["excluded"] >= 1


# =========================================================================== #
#  ENTRY POINT                                                                 #
# =========================================================================== #

if __name__ == "__main__":
    pytest.main([__file__, "-v"])
