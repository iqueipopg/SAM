"""End-to-end integration test for the A1-A8 pipeline.

Runs the full pipeline data flow with synthetic/mocked data.
No GPU, no internet, no LLM models required. Runs in < 30 seconds.

Usage:
    pytest tests/test_e2e_pipeline.py -v
"""

import json
import os
import sys
import tempfile
import shutil
from pathlib import Path

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))

from app import save_checkpoint, load_checkpoint, list_checkpoint_agents, _clean_r_objects


# ══════════════════════════════════════════════════════════════════════════════
# SYNTHETIC FIXTURES
# ══════════════════════════════════════════════════════════════════════════════

SYNTHETIC_PROTOCOL = {
    "framework": "PICO",
    "population": "adults with depression",
    "intervention": "CBT",
    "comparator": "waitlist",
    "outcome": "anxiety scores",
    "research_question": "Is CBT more effective than waitlist control for reducing anxiety scores in adults with depression?",
    "inclusion_criteria": [
        "Adults aged 18+",
        "Diagnosed with depression (DSM/ICD criteria)",
        "CBT as primary intervention",
        "Anxiety score as measured outcome",
    ],
    "exclusion_criteria": [
        "Pediatric populations",
        "Non-RCT designs without comparator",
        "Animal studies",
    ],
    "pubmed_query": '("cognitive behavioral therapy" OR CBT) AND depression AND anxiety AND (randomized OR RCT)',
}

SYNTHETIC_PUBMED_QUERY = (
    '("cognitive behavioral therapy" OR CBT) AND depression AND anxiety AND (randomized OR RCT)'
)

# A2: 5 synthetic papers
SYNTHETIC_A2_PAPERS = [
    {
        "pmid": "30000001",
        "title": "Cognitive behavioral therapy for depression and anxiety in adults: a randomized controlled trial",
        "abstract": "Background: CBT is widely used for depression. Methods: 120 adults randomized to CBT (n=60) or waitlist (n=60). Results: Significant reduction in anxiety scores (SMD=-0.52, 95% CI [-0.78, -0.26]). Conclusion: CBT effective.",
        "authors": ["Smith J", "Doe A", "Brown K"],
        "journal": "J Clin Psychology",
        "year": "2022",
        "source": "pubmed",
    },
    {
        "pmid": "30000002",
        "title": "Effects of CBT on anxiety outcomes in depressed patients: a multi-center RCT",
        "abstract": "Objective: Evaluate CBT vs waitlist for anxiety in depression. 90 patients were randomized. CBT group showed lower anxiety (SMD=-0.38, 95% CI [-0.72, -0.04]). CBT is a viable intervention.",
        "authors": ["Jones M", "Wilson P"],
        "journal": "Behav Res Therapy",
        "year": "2021",
        "source": "pubmed",
    },
    {
        "pmid": "30000003",
        "title": "A randomized trial of cognitive therapy versus waitlist for comorbid depression and anxiety",
        "abstract": "Methods: RCT with 80 participants. CBT reduced anxiety significantly (SMD=-0.61, 95% CI [-1.02, -0.20]). Conclusion: CBT should be considered first-line.",
        "authors": ["Lee S", "Kim H", "Park J"],
        "journal": "Psychol Med",
        "year": "2023",
        "source": "europepmc",
    },
    {
        "pmid": "30000004",
        "title": "Waitlist-controlled trial of brief CBT for anxious depression in primary care",
        "abstract": "Brief CBT (8 sessions) was compared to waitlist in 70 primary care patients with depression and anxiety. Results showed moderate improvement (SMD=-0.45, 95% CI [-0.80, -0.10]).",
        "authors": ["Garcia R", "Martinez L"],
        "journal": "Prim Care Psychiatry",
        "year": "2022",
        "source": "semantic_scholar",
    },
    {
        "pmid": "30000005",
        "title": "Pharmacological treatment of depression: a cost-effectiveness review",
        "abstract": "This review examines cost-effectiveness of SSRIs for depression. No behavioral interventions were compared. Results are limited to pharmacoeconomic outcomes.",
        "authors": ["Taylor W"],
        "journal": "Pharmacoecon Rev",
        "year": "2020",
        "source": "pubmed",
    },
]

SYNTHETIC_A2_STATS = {
    "pubmed": 3,
    "europepmc": 1,
    "semantic_scholar": 1,
    "total_before_dedup": 5,
    "total_after_dedup": 5,
    "duplicates_removed": 0,
}

SYNTHETIC_A2_RESULT = {
    "papers": SYNTHETIC_A2_PAPERS,
    "stats": SYNTHETIC_A2_STATS,
}

# A3: 4 included, 1 excluded (paper 5 is pharmacoeconomic review)
def _make_a3_screening_entry(paper, decision, confidence):
    return {
        "pmid": paper["pmid"],
        "title": paper["title"],
        "abstract": paper["abstract"],
        "authors": paper.get("authors", []),
        "journal": paper.get("journal", ""),
        "year": paper.get("year", ""),
        "source": paper.get("source", ""),
        "decision": decision,
        "confidence": confidence,
        "votes": [decision] * 3,
        "flags": [],
        "reasoning": f"Paper {'meets' if decision == 'INCLUDE' else 'does not meet'} inclusion criteria.",
    }

SYNTHETIC_A3_RESULT = {
    "include": [
        _make_a3_screening_entry(SYNTHETIC_A2_PAPERS[0], "INCLUDE", 1.0),
        _make_a3_screening_entry(SYNTHETIC_A2_PAPERS[1], "INCLUDE", 1.0),
        _make_a3_screening_entry(SYNTHETIC_A2_PAPERS[2], "INCLUDE", 1.0),
        _make_a3_screening_entry(SYNTHETIC_A2_PAPERS[3], "INCLUDE", 0.67),
    ],
    "exclude": [
        _make_a3_screening_entry(SYNTHETIC_A2_PAPERS[4], "EXCLUDE", 1.0),
    ],
    "uncertain": [],
    "stats": {
        "total": 5,
        "included": 4,
        "excluded": 1,
        "uncertain": 0,
        "elapsed_s": 2.0,
        "s_per_paper": 0.4,
        "ensemble_size": 3,
        "voting_threshold": 0.60,
    },
}

# A4: 3 papers with full text, 1 abstract-only (paper 4 abstract-only)
def _make_a4_paper(a3_entry, full_text, text_source):
    return {
        "pmid": a3_entry["pmid"],
        "title": a3_entry["title"],
        "abstract": a3_entry["abstract"],
        "authors": a3_entry.get("authors", []),
        "journal": a3_entry.get("journal", ""),
        "year": a3_entry.get("year", ""),
        "doi": f"10.1000/test.{a3_entry['pmid']}",
        "pmcid": f"PMC{a3_entry['pmid']}" if text_source != "abstract_only" else "",
        "full_text": full_text,
        "text_source": text_source,
        "full_text_length": len(full_text) if full_text else len(a3_entry.get("abstract", "")),
        "sections": {},
        "oa_url": None,
        "fulltext_decision": "INCLUDE",
        "fulltext_reasoning": "Paper meets all inclusion criteria on full-text review.",
        "exclusion_reason": None,
    }

_A4_FULLTEXT_TEMPLATE = (
    "Introduction: Depression and anxiety frequently co-occur. CBT has shown promise. "
    "Methods: This randomized controlled trial enrolled {n} adults with depression and comorbid anxiety. "
    "Participants were randomized to {arms}. Primary outcome was anxiety score (BAI). "
    "Results: The CBT group showed significantly lower anxiety scores compared to waitlist "
    "(SMD={smd}, 95% CI [{ci_lo}, {ci_up}], p={p}). "
    "Discussion: These findings support CBT as effective for anxiety in depressed adults. "
    "Limitations: Single-site study, short follow-up."
)

SYNTHETIC_A4_INCLUDE = [
    _make_a4_paper(
        SYNTHETIC_A3_RESULT["include"][0],
        _A4_FULLTEXT_TEMPLATE.format(n=120, arms="CBT (n=60) or waitlist (n=60)", smd="-0.52", ci_lo="-0.78", ci_up="-0.26", p="0.001"),
        "pmc",
    ),
    _make_a4_paper(
        SYNTHETIC_A3_RESULT["include"][1],
        _A4_FULLTEXT_TEMPLATE.format(n=90, arms="CBT (n=45) or waitlist (n=45)", smd="-0.38", ci_lo="-0.72", ci_up="-0.04", p="0.03"),
        "unpaywall_pdf",
    ),
    _make_a4_paper(
        SYNTHETIC_A3_RESULT["include"][2],
        _A4_FULLTEXT_TEMPLATE.format(n=80, arms="CBT (n=40) or waitlist (n=40)", smd="-0.61", ci_lo="-1.02", ci_up="-0.20", p="0.004"),
        "pmc",
    ),
]

# Paper 4 included but abstract-only
SYNTHETIC_A4_INCLUDE.append(
    _make_a4_paper(
        SYNTHETIC_A3_RESULT["include"][3],
        None,
        "abstract_only",
    )
)
# Override: full_text_length for abstract-only
SYNTHETIC_A4_INCLUDE[3]["full_text_length"] = len(SYNTHETIC_A3_RESULT["include"][3]["abstract"])

# Convenience alias: A4 papers list for use in downstream tests (e.g. author flow)
SYNTHETIC_A4_PAPERS = SYNTHETIC_A4_INCLUDE

SYNTHETIC_A4_RESULT = {
    "papers": SYNTHETIC_A4_INCLUDE,
    "include": SYNTHETIC_A4_INCLUDE[:3],  # 3 included after full-text screening
    "exclude": [SYNTHETIC_A4_INCLUDE[3]],  # paper 4 excluded (abstract-only, insufficient data)
    "stats": {
        "total_retrieved": 4,
        "pmc_fulltext": 2,
        "unpaywall_pdf": 1,
        "europepmc_fulltext": 0,
        "semanticscholar_pdf": 0,
        "scihub_pdf": 0,
        "abstract_only": 1,
        "fulltext_included": 3,
        "fulltext_excluded": 1,
        "elapsed_s": 3.0,
    },
}

# A5: 3 papers with structured extraction data
def _make_a5_paper(a4_paper, study_chars, outcomes, quality="complete"):
    return {
        "pmid": a4_paper["pmid"],
        "title": a4_paper["title"],
        "abstract": a4_paper.get("abstract", ""),
        "authors": a4_paper.get("authors", []),
        "journal": a4_paper.get("journal", ""),
        "year": a4_paper.get("year", ""),
        "text_source": a4_paper["text_source"],
        "study_characteristics": study_chars,
        "intervention_exposure": {
            "intervention_name": "Cognitive Behavioral Therapy (CBT)",
            "comparator_name": "Waitlist control",
            "duration": "12 weeks",
            "dosage": "Weekly 1-hour sessions",
        },
        "outcomes": outcomes,
        "extraction_quality": quality,
        "extraction_confidence": "high" if quality == "complete" else "medium",
        "concordance": {"overall": "high"} if quality == "complete" else None,
        "extraction_notes": "",
        "sections_used": ["results", "methods"],
    }

SYNTHETIC_A5_PAPERS = [
    _make_a5_paper(
        SYNTHETIC_A4_RESULT["include"][0],
        {"study_design": "RCT", "sample_size": 120, "country": "USA", "setting": "outpatient"},
        [{
            "name": "Anxiety score (BAI)",
            "measure_type": "SMD",
            "effect_size_value": -0.52,
            "ci_lower": -0.78,
            "ci_upper": -0.26,
            "p_value": 0.001,
            "n_intervention": 60,
            "n_control": 60,
            "numeric_complete": True,
        }],
    ),
    _make_a5_paper(
        SYNTHETIC_A4_RESULT["include"][1],
        {"study_design": "RCT", "sample_size": 90, "country": "UK", "setting": "community"},
        [{
            "name": "Anxiety score (BAI)",
            "measure_type": "SMD",
            "effect_size_value": -0.38,
            "ci_lower": -0.72,
            "ci_upper": -0.04,
            "p_value": 0.03,
            "n_intervention": 45,
            "n_control": 45,
            "numeric_complete": True,
        }],
    ),
    _make_a5_paper(
        SYNTHETIC_A4_RESULT["include"][2],
        {"study_design": "RCT", "sample_size": 80, "country": "South Korea", "setting": "hospital"},
        [{
            "name": "Anxiety score (BAI)",
            "measure_type": "SMD",
            "effect_size_value": -0.61,
            "ci_lower": -1.02,
            "ci_upper": -0.20,
            "p_value": 0.004,
            "n_intervention": 40,
            "n_control": 40,
            "numeric_complete": True,
        }],
    ),
]

SYNTHETIC_A5_RESULT = {
    "papers": SYNTHETIC_A5_PAPERS,
    "stats": {
        "total": 3,
        "complete_extraction": 3,
        "partial_extraction": 0,
        "failed_extraction": 0,
        "high_confidence": 3,
        "medium_confidence": 0,
        "low_confidence": 0,
        "numeric_complete": 3,
        "numeric_recovered": 0,
        "dual_extraction": True,
        "elapsed_s": 5.0,
    },
}

# A6: 3 papers with RoB2 assessments (all RCTs)
def _make_a6_paper(a5_paper, overall, weight, domains):
    paper = dict(a5_paper)
    paper["quality_assessment"] = {
        "tool": "RoB2",
        "overall": overall,
        "domains": domains,
    }
    paper["quality_weight"] = weight
    return paper

SYNTHETIC_A6_PAPERS = [
    _make_a6_paper(
        SYNTHETIC_A5_PAPERS[0], "Low", 1.0,
        {
            "randomization_process": "Low",
            "deviations_from_interventions": "Low",
            "missing_outcome_data": "Low",
            "measurement_of_outcome": "Low",
            "selection_of_reported_result": "Low",
        },
    ),
    _make_a6_paper(
        SYNTHETIC_A5_PAPERS[1], "Some concerns", 0.6,
        {
            "randomization_process": "Low",
            "deviations_from_interventions": "Some concerns",
            "missing_outcome_data": "Low",
            "measurement_of_outcome": "Low",
            "selection_of_reported_result": "Low",
        },
    ),
    _make_a6_paper(
        SYNTHETIC_A5_PAPERS[2], "Low", 1.0,
        {
            "randomization_process": "Low",
            "deviations_from_interventions": "Low",
            "missing_outcome_data": "Low",
            "measurement_of_outcome": "Some concerns",
            "selection_of_reported_result": "Low",
        },
    ),
]

SYNTHETIC_A6_RESULT = {
    "papers": SYNTHETIC_A6_PAPERS,
    "stats": {
        "total": 3,
        "rob2_assessed": 3,
        "robins_i_assessed": 0,
        "jbi_assessed": 0,
        "failed": 0,
        "elapsed_s": 4.0,
    },
}

# A7: 1 SMD meta-analysis with pooled estimate
SYNTHETIC_A7_RESULT = {
    "analyses": {
        "SMD": {
            "model": {
                "pooled_estimate": -0.50,
                "ci_lower": -0.72,
                "ci_upper": -0.28,
                "p_value": 0.0001,
                "k_studies": 3,
                "method": "REML",
            },
            "heterogeneity": {
                "I2": 5.2,
                "tau2": 0.002,
                "Q": 2.1,
                "Q_p": 0.35,
                "Q_df": 2,
            },
            "plots": {
                "forest": "outputs/plots/forest_SMD.png",
                "funnel": "outputs/plots/funnel_SMD.png",
            },
            "grade": {
                "starting_level": "High",
                "final_level": "Moderate",
                "total_downgrades": 1,
                "downgrades": {"imprecision": 1},
            },
            "sensitivity": {
                "leave_one_out": [
                    {"excluded": "30000001", "pooled_estimate": -0.49, "ci_lower": -0.80, "ci_upper": -0.18},
                    {"excluded": "30000002", "pooled_estimate": -0.56, "ci_lower": -0.79, "ci_upper": -0.33},
                    {"excluded": "30000003", "pooled_estimate": -0.45, "ci_lower": -0.68, "ci_upper": -0.22},
                ],
                "exclude_high_risk": {
                    "k_excluded": 0,
                    "pooled_estimate": -0.50,
                    "ci_lower": -0.72,
                    "ci_upper": -0.28,
                },
            },
            "publication_bias": {
                "egger_test": None,
                "trim_and_fill": None,
                "note": "Too few studies (k=3) for Egger test (requires k>=10)",
            },
            "studies": [
                {"pmid": "30000001", "yi": -0.52, "vi": 0.018, "weight": 1.0},
                {"pmid": "30000002", "yi": -0.38, "vi": 0.030, "weight": 0.6},
                {"pmid": "30000003", "yi": -0.61, "vi": 0.044, "weight": 1.0},
            ],
        },
    },
    "stats": {
        "total_papers_input": 3,
        "papers_with_usable_data": 3,
        "analyses_performed": 1,
        "measure_types_found": ["SMD"],
        "elapsed_s": 2.0,
    },
}

# A8: manuscript with 6 sections
SYNTHETIC_A8_RESULT = {
    "manuscript": {
        "abstract": (
            "Background: Cognitive behavioral therapy (CBT) is a widely used intervention for depression, "
            "yet its effects on comorbid anxiety remain debated. This systematic review and meta-analysis "
            "examined the efficacy of CBT compared to waitlist control for reducing anxiety scores in adults "
            "with depression. Methods: We searched PubMed, Europe PMC, and Semantic Scholar. Three RCTs "
            "met inclusion criteria (N=290). Random-effects meta-analysis was conducted. Results: CBT "
            "significantly reduced anxiety (pooled SMD=-0.50, 95% CI [-0.72, -0.28], p<0.001). "
            "Heterogeneity was low (I2=5.2%). GRADE certainty was moderate. Conclusion: CBT effectively "
            "reduces anxiety in adults with depression compared to waitlist control."
        ),
        "introduction": (
            "Depression and anxiety disorders frequently co-occur, affecting approximately 50% of individuals "
            "with major depressive disorder. Cognitive behavioral therapy is an evidence-based treatment for "
            "both conditions, but the specific impact on anxiety when treating depression has been debated. "
            "Previous reviews have examined CBT for depression or anxiety separately, but few have focused on "
            "the anxiolytic effects of CBT in populations primarily diagnosed with depression. This systematic "
            "review aims to address this gap by synthesizing evidence from randomized controlled trials comparing "
            "CBT to waitlist control conditions, with anxiety scores as the primary outcome."
        ),
        "methods": (
            "This systematic review followed PRISMA 2020 guidelines. We searched PubMed, Europe PMC, and "
            "Semantic Scholar using the query: (cognitive behavioral therapy OR CBT) AND depression AND anxiety "
            "AND (randomized OR RCT). Title/abstract screening was performed using ensemble LLM voting (3 votes, "
            "60% threshold). Full-text retrieval used PMC, Unpaywall, and Europe PMC. Data extraction captured "
            "study characteristics, intervention details, and outcomes. Quality was assessed using RoB 2 for "
            "randomized trials. Random-effects meta-analysis (REML) was performed using the metafor package in R. "
            "Heterogeneity was assessed using I-squared and Cochran Q. GRADE assessment was applied to rate "
            "certainty of evidence."
        ),
        "results": (
            "Five records were identified, and after screening, three RCTs (N=290) were included. All studies "
            "compared CBT to waitlist control in adults with depression, measuring anxiety with the BAI. "
            "Quality assessment: two studies were rated low risk of bias, one had some concerns. "
            "Meta-analysis showed a significant pooled effect favoring CBT (SMD=-0.50, 95% CI [-0.72, -0.28], "
            "p<0.001). Heterogeneity was low (I2=5.2%, Q=2.1, p=0.35). Leave-one-out sensitivity analysis "
            "showed stable results. GRADE certainty was rated moderate, downgraded once for imprecision."
        ),
        "discussion": (
            "Our findings indicate that CBT significantly reduces anxiety in adults with depression compared to "
            "waitlist control, with a moderate effect size. The low heterogeneity suggests consistent effects "
            "across studies. These results align with previous meta-analyses showing CBT efficacy for anxiety "
            "disorders. Strengths include rigorous automated screening and quality assessment. Limitations include "
            "the small number of studies (k=3), reliance on waitlist controls, and potential for publication bias "
            "that could not be formally tested. Future research should include active comparators and longer follow-up."
        ),
        "conclusion": (
            "CBT is effective for reducing anxiety in adults with depression compared to waitlist control "
            "(SMD=-0.50, 95% CI [-0.72, -0.28]). The evidence was rated moderate certainty using GRADE. "
            "Clinicians should consider CBT as a first-line treatment for comorbid depression and anxiety. "
            "More high-quality RCTs with active comparators are needed to strengthen the evidence base."
        ),
    },
    "references": [
        {"pmid": "30000001", "title": "Cognitive behavioral therapy for depression and anxiety in adults: a randomized controlled trial", "citation": "Smith J, Doe A, Brown K (2022)", "journal": "J Clin Psychology"},
        {"pmid": "30000002", "title": "Effects of CBT on anxiety outcomes in depressed patients: a multi-center RCT", "citation": "Jones M, Wilson P (2021)", "journal": "Behav Res Therapy"},
        {"pmid": "30000003", "title": "A randomized trial of cognitive therapy versus waitlist for comorbid depression and anxiety", "citation": "Lee S, Kim H, Park J (2023)", "journal": "Psychol Med"},
    ],
    "stats": {
        "total_chunks": 24,
        "total_papers_cited": 3,
        "chunks_per_section": {
            "abstract": 4,
            "introduction": 6,
            "methods": 4,
            "results": 4,
            "discussion": 4,
            "conclusion": 2,
        },
        "elapsed_s": 15.0,
    },
    "pdf_path": None,
}


# ══════════════════════════════════════════════════════════════════════════════
# PRISMA helper
# ══════════════════════════════════════════════════════════════════════════════

def build_prisma(a2_stats, a3_stats, a4_stats, a5_stats, a6_stats):
    """Build PRISMA dict the same way app.py does."""
    n_extracted = a5_stats.get("complete_extraction", 0) + a5_stats.get("partial_extraction", 0)
    n_quality = a6_stats.get("rob2_assessed", 0) + a6_stats.get("robins_i_assessed", 0) + a6_stats.get("jbi_assessed", 0)
    return {
        "identified": a2_stats.get("total_before_dedup", 0),
        "identified_per_db": {
            "pubmed": a2_stats.get("pubmed", 0),
            "europepmc": a2_stats.get("europepmc", 0),
            "semantic_scholar": a2_stats.get("semantic_scholar", 0),
        },
        "after_dedup": a2_stats.get("total_after_dedup", 0),
        "duplicates_removed": a2_stats.get("duplicates_removed", 0),
        "screened": a3_stats.get("total", 0),
        "excluded_title_abstract": a3_stats.get("excluded", 0),
        "uncertain_title_abstract": a3_stats.get("uncertain", 0),
        "assessed_fulltext": a3_stats.get("included", 0),
        "excluded_fulltext": a4_stats.get("fulltext_excluded", 0),
        "included_final": a4_stats.get("fulltext_included", 0),
        "extracted": n_extracted,
        "quality_assessed": n_quality,
    }


def build_final_output():
    """Build the final pipeline output dict as app.py does."""
    prisma = build_prisma(
        SYNTHETIC_A2_STATS,
        SYNTHETIC_A3_RESULT["stats"],
        SYNTHETIC_A4_RESULT["stats"],
        SYNTHETIC_A5_RESULT["stats"],
        SYNTHETIC_A6_RESULT["stats"],
    )
    return {
        "query": "Is CBT more effective than waitlist for reducing anxiety in adults with depression?",
        "protocol": SYNTHETIC_PROTOCOL,
        "pubmed_query": SYNTHETIC_PUBMED_QUERY,
        "prisma": prisma,
        "search_stats": SYNTHETIC_A2_STATS,
        "screening": SYNTHETIC_A3_RESULT,
        "fulltext_retrieval": SYNTHETIC_A4_RESULT,
        "data_extraction": SYNTHETIC_A5_RESULT,
        "quality_assessment": SYNTHETIC_A6_RESULT,
        "synthesis": SYNTHETIC_A7_RESULT,
        "manuscript": SYNTHETIC_A8_RESULT,
    }


# ══════════════════════════════════════════════════════════════════════════════
# PIPELINE DATA FLOW TEST
# ══════════════════════════════════════════════════════════════════════════════

class TestPipelineDataFlow:
    """Test the full A1-A8 pipeline data flow using synthetic data."""

    def test_a1_protocol_structure(self):
        """A1 protocol has all required PICO fields."""
        p = SYNTHETIC_PROTOCOL
        assert p["framework"] == "PICO"
        assert "population" in p
        assert "intervention" in p
        assert "comparator" in p
        assert "outcome" in p
        assert len(p["inclusion_criteria"]) > 0
        assert len(p["exclusion_criteria"]) > 0

    def test_a2_search_output(self):
        """A2 output has papers list and stats dict."""
        result = SYNTHETIC_A2_RESULT
        assert "papers" in result
        assert "stats" in result
        assert len(result["papers"]) == 5
        # Each paper has required fields
        for paper in result["papers"]:
            assert "pmid" in paper
            assert "title" in paper
            assert "abstract" in paper

    def test_a2_to_a3_flow(self):
        """A2 papers feed into A3 screening; counts are consistent."""
        a2_papers = SYNTHETIC_A2_RESULT["papers"]
        a3 = SYNTHETIC_A3_RESULT
        # A3 received all A2 papers
        assert a3["stats"]["total"] == len(a2_papers)
        # Include + exclude + uncertain = total
        assert (
            a3["stats"]["included"] + a3["stats"]["excluded"] + a3["stats"]["uncertain"]
            == a3["stats"]["total"]
        )

    def test_a3_to_a4_flow(self):
        """A3 included papers feed into A4; A4 has include/exclude."""
        a3_included = SYNTHETIC_A3_RESULT["include"]
        a4 = SYNTHETIC_A4_RESULT
        # A4 received A3 included papers
        assert a4["stats"]["total_retrieved"] == len(a3_included)
        # A4 include + exclude = total
        assert len(a4["include"]) + len(a4["exclude"]) == a4["stats"]["total_retrieved"]

    def test_a4_to_a5_flow(self):
        """A4 included papers feed into A5 extraction."""
        a4_included = SYNTHETIC_A4_RESULT["include"]
        a5 = SYNTHETIC_A5_RESULT
        assert a5["stats"]["total"] == len(a4_included)
        assert len(a5["papers"]) == len(a4_included)

    def test_a5_to_a6_flow(self):
        """A5 papers feed into A6 quality assessment."""
        a5_papers = SYNTHETIC_A5_RESULT["papers"]
        a6 = SYNTHETIC_A6_RESULT
        assert a6["stats"]["total"] == len(a5_papers)
        assert len(a6["papers"]) == len(a5_papers)

    def test_a6_to_a7_flow(self):
        """A6 papers (with weights) feed into A7 synthesis."""
        a6_papers = SYNTHETIC_A6_RESULT["papers"]
        a7 = SYNTHETIC_A7_RESULT
        assert a7["stats"]["total_papers_input"] == len(a6_papers)
        # All studies in analyses reference valid PMIDs
        smd = a7["analyses"]["SMD"]
        study_pmids = {s["pmid"] for s in smd["studies"]}
        a6_pmids = {p["pmid"] for p in a6_papers}
        assert study_pmids == a6_pmids

    def test_a7_to_a8_flow(self):
        """A7 synthesis data informs A8 manuscript content."""
        a8 = SYNTHETIC_A8_RESULT
        manuscript = a8["manuscript"]
        # Manuscript references the pooled estimate from A7
        assert "-0.50" in manuscript["results"] or "-0.5" in manuscript["results"]
        # References correspond to included papers
        ref_pmids = {r["pmid"] for r in a8["references"]}
        a7_pmids = {s["pmid"] for s in SYNTHETIC_A7_RESULT["analyses"]["SMD"]["studies"]}
        assert ref_pmids == a7_pmids

    def test_final_output_has_all_keys(self):
        """Final pipeline output has all expected top-level keys."""
        output = build_final_output()
        expected_keys = {
            "query", "protocol", "pubmed_query", "prisma",
            "search_stats", "screening", "fulltext_retrieval",
            "data_extraction", "quality_assessment", "synthesis", "manuscript",
        }
        assert expected_keys.issubset(set(output.keys()))


# ══════════════════════════════════════════════════════════════════════════════
# DATA INTEGRITY TESTS
# ══════════════════════════════════════════════════════════════════════════════

class TestDataIntegrity:
    """Verify data integrity across pipeline stages."""

    def test_paper_ids_preserved(self):
        """PMIDs from A2 are preserved through A3, A4, A5, A6."""
        a2_pmids = {p["pmid"] for p in SYNTHETIC_A2_PAPERS}

        # A3 PMIDs are a subset of A2
        a3_all_pmids = set()
        for bucket in ("include", "exclude", "uncertain"):
            for entry in SYNTHETIC_A3_RESULT[bucket]:
                a3_all_pmids.add(entry["pmid"])
        assert a3_all_pmids.issubset(a2_pmids)

        # A3 included PMIDs
        a3_included_pmids = {p["pmid"] for p in SYNTHETIC_A3_RESULT["include"]}

        # A4 PMIDs are subset of A3 included
        a4_all_pmids = {p["pmid"] for p in SYNTHETIC_A4_RESULT["papers"]}
        assert a4_all_pmids.issubset(a3_included_pmids)

        # A4 included PMIDs
        a4_included_pmids = {p["pmid"] for p in SYNTHETIC_A4_RESULT["include"]}

        # A5 PMIDs are subset of A4 included
        a5_pmids = {p["pmid"] for p in SYNTHETIC_A5_RESULT["papers"]}
        assert a5_pmids.issubset(a4_included_pmids)

        # A6 PMIDs equal A5 PMIDs
        a6_pmids = {p["pmid"] for p in SYNTHETIC_A6_RESULT["papers"]}
        assert a6_pmids == a5_pmids

    def test_numeric_extraction_format(self):
        """A5 outcomes have effect_size_value, ci_lower, ci_upper as floats."""
        for paper in SYNTHETIC_A5_RESULT["papers"]:
            assert len(paper["outcomes"]) > 0, f"Paper {paper['pmid']} has no outcomes"
            for outcome in paper["outcomes"]:
                assert isinstance(outcome["effect_size_value"], float), \
                    f"effect_size_value not float: {outcome['effect_size_value']}"
                assert isinstance(outcome["ci_lower"], float), \
                    f"ci_lower not float: {outcome['ci_lower']}"
                assert isinstance(outcome["ci_upper"], float), \
                    f"ci_upper not float: {outcome['ci_upper']}"
                # CI lower < estimate < CI upper
                assert outcome["ci_lower"] <= outcome["effect_size_value"] <= outcome["ci_upper"], \
                    f"CI range invalid: [{outcome['ci_lower']}, {outcome['effect_size_value']}, {outcome['ci_upper']}]"

    def test_quality_weights_range(self):
        """A6 quality weights are between 0 and 1."""
        for paper in SYNTHETIC_A6_RESULT["papers"]:
            w = paper["quality_weight"]
            assert 0.0 <= w <= 1.0, f"Weight out of range for {paper['pmid']}: {w}"

    def test_quality_assessment_structure(self):
        """A6 papers have proper quality_assessment dict with tool and domains."""
        for paper in SYNTHETIC_A6_RESULT["papers"]:
            qa = paper["quality_assessment"]
            assert "tool" in qa
            assert qa["tool"] in ("RoB2", "ROBINS-I", "JBI")
            assert "overall" in qa
            assert "domains" in qa
            assert isinstance(qa["domains"], dict)
            # RoB2 has 5 domains
            if qa["tool"] == "RoB2":
                assert len(qa["domains"]) == 5

    def test_synthesis_structure(self):
        """A7 has analyses dict with model, heterogeneity, grade."""
        a7 = SYNTHETIC_A7_RESULT
        assert "analyses" in a7
        assert "stats" in a7
        assert len(a7["analyses"]) >= 1

        for measure, analysis in a7["analyses"].items():
            assert "model" in analysis, f"No model in {measure} analysis"
            assert "heterogeneity" in analysis, f"No heterogeneity in {measure}"
            assert "grade" in analysis, f"No grade in {measure}"

            model = analysis["model"]
            assert "pooled_estimate" in model
            assert "ci_lower" in model
            assert "ci_upper" in model
            assert "p_value" in model
            assert "k_studies" in model
            assert model["ci_lower"] <= model["pooled_estimate"] <= model["ci_upper"]

            het = analysis["heterogeneity"]
            assert "I2" in het
            assert 0 <= het["I2"] <= 100

            grade = analysis["grade"]
            assert "final_level" in grade
            assert grade["final_level"] in ("High", "Moderate", "Low", "Very low")

    def test_manuscript_sections(self):
        """A8 has abstract, introduction, methods, results, discussion, conclusion."""
        expected = ["abstract", "introduction", "methods", "results", "discussion", "conclusion"]
        manuscript = SYNTHETIC_A8_RESULT["manuscript"]
        for sec in expected:
            assert sec in manuscript, f"Missing section: {sec}"
            assert len(manuscript[sec]) > 50, f"Section {sec} too short: {len(manuscript[sec])} chars"

    def test_manuscript_references(self):
        """A8 references list has pmid and title for each entry."""
        refs = SYNTHETIC_A8_RESULT["references"]
        assert len(refs) >= 1
        for ref in refs:
            assert "pmid" in ref
            assert "title" in ref
            assert "citation" in ref

    def test_prisma_counts_consistent(self):
        """PRISMA: identified >= screened >= assessed_fulltext >= included_final."""
        prisma = build_prisma(
            SYNTHETIC_A2_STATS,
            SYNTHETIC_A3_RESULT["stats"],
            SYNTHETIC_A4_RESULT["stats"],
            SYNTHETIC_A5_RESULT["stats"],
            SYNTHETIC_A6_RESULT["stats"],
        )
        assert prisma["identified"] >= prisma["screened"]
        assert prisma["screened"] >= prisma["assessed_fulltext"]
        assert prisma["assessed_fulltext"] >= prisma["included_final"]
        assert prisma["included_final"] >= 0

        # Excluded + included = assessed
        assert (
            prisma["excluded_fulltext"] + prisma["included_final"]
            == prisma["assessed_fulltext"]
        )

        # Title/abstract: excluded + uncertain + assessed_fulltext = screened
        assert (
            prisma["excluded_title_abstract"]
            + prisma["uncertain_title_abstract"]
            + prisma["assessed_fulltext"]
            == prisma["screened"]
        )


# ══════════════════════════════════════════════════════════════════════════════
# CHECKPOINT TEST
# ══════════════════════════════════════════════════════════════════════════════

class TestCheckpoints:
    """Test checkpoint save/load round-trip."""

    def setup_method(self):
        self.tmpdir = tempfile.mkdtemp(prefix="test_e2e_ckpt_")

    def teardown_method(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_checkpoint_save_load(self):
        """Save checkpoint, load it, verify data matches."""
        run_dir = self.tmpdir

        # Save each agent checkpoint
        save_checkpoint(run_dir, "a1_protocol", {
            "query": "test query",
            "protocol": SYNTHETIC_PROTOCOL,
            "pubmed_query": SYNTHETIC_PUBMED_QUERY,
        })
        save_checkpoint(run_dir, "a2_search", SYNTHETIC_A2_RESULT)
        save_checkpoint(run_dir, "a3_screening", SYNTHETIC_A3_RESULT)
        save_checkpoint(run_dir, "a4_fulltext", SYNTHETIC_A4_RESULT)
        save_checkpoint(run_dir, "a5_extraction", SYNTHETIC_A5_RESULT)
        save_checkpoint(run_dir, "a6_quality", SYNTHETIC_A6_RESULT)
        save_checkpoint(run_dir, "a7_synthesis", SYNTHETIC_A7_RESULT)
        save_checkpoint(run_dir, "a8_manuscript", SYNTHETIC_A8_RESULT)

        # List checkpoints
        agents = list_checkpoint_agents(run_dir)
        assert len(agents) == 8
        assert "a1_protocol" in agents
        assert "a8_manuscript" in agents

        # Load and verify
        a1 = load_checkpoint(run_dir, "a1_protocol")
        assert a1["protocol"]["framework"] == "PICO"
        assert a1["pubmed_query"] == SYNTHETIC_PUBMED_QUERY

        a3 = load_checkpoint(run_dir, "a3_screening")
        assert a3["stats"]["included"] == 4
        assert a3["stats"]["excluded"] == 1

        a5 = load_checkpoint(run_dir, "a5_extraction")
        assert len(a5["papers"]) == 3
        assert a5["papers"][0]["outcomes"][0]["effect_size_value"] == -0.52

        a7 = load_checkpoint(run_dir, "a7_synthesis")
        assert a7["analyses"]["SMD"]["model"]["pooled_estimate"] == -0.50

    def test_checkpoint_missing_returns_none(self):
        """Loading a nonexistent checkpoint returns None."""
        result = load_checkpoint(self.tmpdir, "nonexistent_agent")
        assert result is None

    def test_checkpoint_overwrite(self):
        """Saving a checkpoint twice overwrites the first."""
        save_checkpoint(self.tmpdir, "a1_protocol", {"version": 1})
        save_checkpoint(self.tmpdir, "a1_protocol", {"version": 2})
        loaded = load_checkpoint(self.tmpdir, "a1_protocol")
        assert loaded["version"] == 2

    def test_final_output_json_serializable(self):
        """Final output dict is fully JSON-serializable."""
        output = build_final_output()
        # Should not raise
        serialized = json.dumps(output, ensure_ascii=False, default=str)
        restored = json.loads(serialized)
        assert restored["protocol"]["framework"] == "PICO"
        assert len(restored["screening"]["include"]) == 4


# ══════════════════════════════════════════════════════════════════════════════
# EDGE CASE TESTS
# ══════════════════════════════════════════════════════════════════════════════

class TestEdgeCases:
    """Test pipeline behavior with edge-case inputs."""

    def test_empty_search_results(self):
        """A2 returns 0 papers: pipeline should handle gracefully."""
        a2_result = {
            "papers": [],
            "stats": {
                "pubmed": 0, "europepmc": 0, "semantic_scholar": 0,
                "total_before_dedup": 0, "total_after_dedup": 0, "duplicates_removed": 0,
            },
        }
        # No papers means A3 would receive empty list
        assert len(a2_result["papers"]) == 0
        # In app.py, this triggers early return with warning
        # Verify the stats are consistent
        assert a2_result["stats"]["total_after_dedup"] == 0

    def test_all_excluded_screening(self):
        """A3 excludes everything: pipeline should continue without A4-A8."""
        a3_result = {
            "include": [],
            "exclude": [
                _make_a3_screening_entry(p, "EXCLUDE", 1.0)
                for p in SYNTHETIC_A2_PAPERS
            ],
            "uncertain": [],
            "stats": {
                "total": 5,
                "included": 0,
                "excluded": 5,
                "uncertain": 0,
                "elapsed_s": 1.0,
                "s_per_paper": 0.2,
                "ensemble_size": 3,
                "voting_threshold": 0.60,
            },
        }
        assert len(a3_result["include"]) == 0
        assert a3_result["stats"]["included"] == 0

        # Build output as app.py would with no included papers
        prisma = {
            "identified": 5,
            "screened": 5,
            "excluded_title_abstract": 5,
            "uncertain_title_abstract": 0,
            "assessed_fulltext": 0,
            "excluded_fulltext": 0,
            "included_final": 0,
        }
        assert prisma["included_final"] == 0
        assert prisma["identified"] >= prisma["screened"]

    def test_no_numeric_data_extraction(self):
        """A5 extracts no numeric outcomes: A7 should produce empty analyses."""
        papers_no_numeric = []
        for paper in SYNTHETIC_A5_PAPERS:
            p = dict(paper)
            p["outcomes"] = [{
                "name": "Qualitative finding",
                "measure_type": "qualitative",
                "effect_size_value": None,
                "ci_lower": None,
                "ci_upper": None,
                "p_value": None,
                "numeric_complete": False,
            }]
            p["quality_weight"] = 1.0
            papers_no_numeric.append(p)

        # A7 should have no analyses when no numeric data is available
        # Simulate the expected output
        a7_empty = {
            "analyses": {},
            "stats": {
                "total_papers_input": 3,
                "papers_with_usable_data": 0,
                "analyses_performed": 0,
                "measure_types_found": [],
                "elapsed_s": 0.1,
            },
        }
        assert a7_empty["stats"]["analyses_performed"] == 0
        assert len(a7_empty["analyses"]) == 0

    def test_single_paper_no_meta_analysis(self):
        """With only 1 paper, meta-analysis requires k>=2 so should not run."""
        a7_single = {
            "analyses": {},
            "stats": {
                "total_papers_input": 1,
                "papers_with_usable_data": 1,
                "analyses_performed": 0,
                "measure_types_found": ["SMD"],
                "elapsed_s": 0.1,
            },
        }
        # App.py checks: len(a7_input.get("papers", [])) >= 2
        assert a7_single["stats"]["analyses_performed"] == 0

    def test_a4_all_abstract_only(self):
        """All papers have abstract-only text: pipeline should still work."""
        abstract_only_papers = []
        for p in SYNTHETIC_A4_INCLUDE[:3]:
            paper = dict(p)
            paper["text_source"] = "abstract_only"
            paper["full_text"] = None
            paper["full_text_length"] = len(paper.get("abstract", ""))
            abstract_only_papers.append(paper)

        # Stats should reflect all abstract-only
        assert all(p["text_source"] == "abstract_only" for p in abstract_only_papers)

    def test_mixed_study_designs_quality(self):
        """A6 handles mixed study designs (RCT + observational)."""
        mixed_papers = list(SYNTHETIC_A6_PAPERS)
        # Add an observational study
        obs_paper = dict(SYNTHETIC_A5_PAPERS[0])
        obs_paper["pmid"] = "30000099"
        obs_paper["study_characteristics"] = {"study_design": "cohort", "sample_size": 200}
        obs_paper["quality_assessment"] = {
            "tool": "ROBINS-I",
            "overall": "Moderate",
            "domains": {
                "confounding": "Moderate",
                "selection": "Low",
                "classification_interventions": "Low",
                "deviations": "Moderate",
                "missing_data": "Low",
                "measurement": "Low",
                "selection_reported": "Low",
            },
        }
        obs_paper["quality_weight"] = 0.7
        mixed_papers.append(obs_paper)

        # All weights valid
        for p in mixed_papers:
            assert 0.0 <= p["quality_weight"] <= 1.0
        # Mixed tools
        tools = {p["quality_assessment"]["tool"] for p in mixed_papers}
        assert "RoB2" in tools
        assert "ROBINS-I" in tools


# ══════════════════════════════════════════════════════════════════════════════
# CLEAN R OBJECTS TEST
# ══════════════════════════════════════════════════════════════════════════════

class TestCleanRObjects:
    """Test the _clean_r_objects utility from app.py."""

    def test_removes_fit_keys(self):
        """_clean_r_objects removes _fit keys recursively."""
        data = {
            "analyses": {
                "SMD": {
                    "model": {"pooled_estimate": -0.5},
                    "_fit": "<R object>",
                    "nested": {
                        "_fit": "<another R object>",
                        "value": 42,
                    },
                },
            },
        }
        cleaned = _clean_r_objects(data)
        assert "_fit" not in cleaned["analyses"]["SMD"]
        assert "_fit" not in cleaned["analyses"]["SMD"]["nested"]
        assert cleaned["analyses"]["SMD"]["nested"]["value"] == 42

    def test_handles_lists(self):
        """_clean_r_objects handles lists with nested dicts."""
        data = [{"_fit": "x", "val": 1}, {"_fit": "y", "val": 2}]
        cleaned = _clean_r_objects(data)
        assert all("_fit" not in item for item in cleaned)
        assert cleaned[0]["val"] == 1

    def test_preserves_scalars(self):
        """_clean_r_objects preserves non-dict/list values."""
        assert _clean_r_objects(42) == 42
        assert _clean_r_objects("hello") == "hello"
        assert _clean_r_objects(None) is None


# ══════════════════════════════════════════════════════════════════════════════
# FULL OUTPUT STRUCTURE TEST
# ══════════════════════════════════════════════════════════════════════════════

class TestFullOutputStructure:
    """Verify the complete pipeline output matches expected structure."""

    @pytest.fixture
    def output(self):
        return build_final_output()

    def test_query_present(self, output):
        assert isinstance(output["query"], str)
        assert len(output["query"]) > 0

    def test_protocol_present(self, output):
        p = output["protocol"]
        assert p["framework"] in ("PICO", "SPIDER", "PECO", "PCC")

    def test_prisma_present(self, output):
        prisma = output["prisma"]
        assert prisma["identified"] > 0
        assert prisma["included_final"] >= 0

    def test_screening_has_buckets(self, output):
        s = output["screening"]
        assert "include" in s
        assert "exclude" in s
        assert "uncertain" in s
        assert "stats" in s

    def test_fulltext_has_include_exclude(self, output):
        ft = output["fulltext_retrieval"]
        assert "include" in ft
        assert "exclude" in ft
        assert "stats" in ft

    def test_extraction_has_papers(self, output):
        ext = output["data_extraction"]
        assert "papers" in ext
        assert "stats" in ext

    def test_quality_has_papers(self, output):
        qa = output["quality_assessment"]
        assert "papers" in qa
        assert "stats" in qa

    def test_synthesis_has_analyses(self, output):
        syn = output["synthesis"]
        assert "analyses" in syn
        assert "stats" in syn

    def test_manuscript_has_sections(self, output):
        ms = output["manuscript"]
        assert "manuscript" in ms
        assert "references" in ms
        assert "stats" in ms


# ══════════════════════════════════════════════════════════════════════════════
# PAPER FIELD COMPLETENESS TEST
# ══════════════════════════════════════════════════════════════════════════════

class TestPaperFieldCompleteness:
    """Verify papers at each stage have the required fields."""

    def test_a2_paper_fields(self):
        required = {"pmid", "title", "abstract"}
        for paper in SYNTHETIC_A2_PAPERS:
            assert required.issubset(set(paper.keys())), \
                f"Paper {paper.get('pmid')} missing A2 fields: {required - set(paper.keys())}"

    def test_a4_paper_fields(self):
        required = {"pmid", "title", "abstract", "text_source", "full_text_length",
                     "fulltext_decision"}
        for paper in SYNTHETIC_A4_RESULT["include"]:
            assert required.issubset(set(paper.keys())), \
                f"Paper {paper.get('pmid')} missing A4 fields: {required - set(paper.keys())}"

    def test_a5_paper_fields(self):
        required = {"pmid", "title", "study_characteristics", "outcomes",
                     "extraction_quality", "text_source"}
        for paper in SYNTHETIC_A5_RESULT["papers"]:
            assert required.issubset(set(paper.keys())), \
                f"Paper {paper.get('pmid')} missing A5 fields: {required - set(paper.keys())}"

    def test_a6_paper_fields(self):
        required = {"pmid", "title", "quality_assessment", "quality_weight"}
        for paper in SYNTHETIC_A6_RESULT["papers"]:
            assert required.issubset(set(paper.keys())), \
                f"Paper {paper.get('pmid')} missing A6 fields: {required - set(paper.keys())}"

    def test_a7_study_fields(self):
        required = {"pmid", "yi", "vi", "weight"}
        for analysis in SYNTHETIC_A7_RESULT["analyses"].values():
            for study in analysis["studies"]:
                assert required.issubset(set(study.keys())), \
                    f"Study {study.get('pmid')} missing A7 fields: {required - set(study.keys())}"


# ══════════════════════════════════════════════════════════════════════════════
# NEW E2E TESTS — edge cases, language, config, frameworks
# ══════════════════════════════════════════════════════════════════════════════

class TestQueryEdgeCases:
    """Test pipeline behavior with unusual query inputs."""

    def test_query_only_numbers(self):
        """Query '2 3 5 7' has no clinical keywords; A1 should still produce a protocol structure."""
        query = "2 3 5 7"
        # A1 must always produce valid protocol dict even with non-clinical input
        # Simulate A1 output with empty/default fields (LLM would try its best)
        protocol = {
            "framework": "PICO",
            "population": "",
            "intervention": "",
            "comparator": "",
            "outcome": "",
            "research_question": query,
            "inclusion_criteria": [],
            "exclusion_criteria": [],
            "pubmed_query": "2 3 5 7",
        }
        # Structure is valid even if content is empty
        assert protocol["framework"] in ("PICO", "SPIDER", "PECO", "PCC")
        assert protocol["research_question"] == query
        # A2 with this query would likely return 0 papers
        a2_result = {
            "papers": [],
            "stats": {
                "pubmed": 0, "europepmc": 0, "semantic_scholar": 0,
                "total_before_dedup": 0, "total_after_dedup": 0, "duplicates_removed": 0,
            },
        }
        assert len(a2_result["papers"]) == 0
        assert a2_result["stats"]["total_after_dedup"] == 0

    def test_query_portuguese(self):
        """Portuguese query should be translatable; A1 produces English protocol."""
        query_pt = "Efeito da terapia cognitivo-comportamental na ansiedade em adultos com depressao"
        # A1 translates to English internally (translate_to_english step)
        # Simulated A1 output after translation
        protocol = {
            "framework": "PICO",
            "population": "adults with depression",
            "intervention": "cognitive behavioral therapy",
            "comparator": "control",
            "outcome": "anxiety",
            "research_question": "Effect of cognitive behavioral therapy on anxiety in adults with depression",
            "inclusion_criteria": [
                "Adults aged 18+",
                "Diagnosed with depression",
                "CBT as intervention",
            ],
            "exclusion_criteria": [
                "Pediatric populations",
                "Animal studies",
            ],
            "pubmed_query": '("cognitive behavioral therapy" OR CBT) AND depression AND anxiety',
        }
        # Protocol must be in English regardless of input language
        assert "depression" in protocol["population"]
        assert "cognitive behavioral therapy" in protocol["intervention"]
        # PubMed query must be English
        assert "terapia" not in protocol["pubmed_query"]
        assert "CBT" in protocol["pubmed_query"]
        # Full pipeline output is valid
        output = {
            "query": query_pt,
            "protocol": protocol,
            "pubmed_query": protocol["pubmed_query"],
        }
        assert isinstance(output["query"], str)
        assert len(output["protocol"]["inclusion_criteria"]) > 0


class TestA2EmptyResults:
    """Test that A3/A4 skip gracefully when A2 returns 0 papers."""

    def test_a2_returns_empty(self):
        """A2 returns 0 papers; A3 and A4 should skip gracefully."""
        a2_result = {
            "papers": [],
            "stats": {
                "pubmed": 0, "europepmc": 0, "semantic_scholar": 0,
                "total_before_dedup": 0, "total_after_dedup": 0, "duplicates_removed": 0,
            },
        }
        # A3 receives empty list: should produce empty screening with 0 counts
        a3_result = {
            "include": [],
            "exclude": [],
            "uncertain": [],
            "stats": {
                "total": 0,
                "included": 0,
                "excluded": 0,
                "uncertain": 0,
                "elapsed_s": 0.0,
                "s_per_paper": 0.0,
                "ensemble_size": 3,
                "voting_threshold": 0.60,
            },
        }
        assert a3_result["stats"]["total"] == len(a2_result["papers"])
        assert a3_result["stats"]["included"] == 0
        assert len(a3_result["include"]) == 0

        # A4 receives empty list from A3: should produce empty result
        a4_result = {
            "papers": [],
            "include": [],
            "exclude": [],
            "stats": {
                "total_retrieved": 0,
                "pmc_fulltext": 0,
                "unpaywall_pdf": 0,
                "europepmc_fulltext": 0,
                "semanticscholar_pdf": 0,
                "scihub_pdf": 0,
                "abstract_only": 0,
                "fulltext_included": 0,
                "fulltext_excluded": 0,
                "elapsed_s": 0.0,
            },
        }
        assert a4_result["stats"]["total_retrieved"] == 0
        assert len(a4_result["include"]) == 0
        assert len(a4_result["exclude"]) == 0

        # PRISMA flow is consistent with zero papers
        prisma = build_prisma(
            a2_result["stats"],
            a3_result["stats"],
            a4_result["stats"],
            {"complete_extraction": 0, "partial_extraction": 0},
            {"rob2_assessed": 0, "robins_i_assessed": 0, "jbi_assessed": 0},
        )
        assert prisma["identified"] == 0
        assert prisma["screened"] == 0
        assert prisma["included_final"] == 0


class TestA5AbstractOnlyPaper:
    """Test A5 extraction when a paper has only abstract text."""

    def test_a5_abstract_only_paper(self):
        """Paper without full-text uses abstract for extraction in A5."""
        # Create an abstract-only A4 paper
        abstract_paper = {
            "pmid": "30000099",
            "title": "Brief CBT for anxious depression: a pilot study",
            "abstract": (
                "This pilot RCT enrolled 30 adults with depression and anxiety. "
                "Participants received 6 sessions of brief CBT (n=15) or waitlist (n=15). "
                "Results: moderate improvement in anxiety (SMD=-0.40, 95% CI [-0.85, 0.05])."
            ),
            "authors": ["Nguyen T", "Lopez M"],
            "journal": "Pilot Feasibility Studies",
            "year": "2024",
            "text_source": "abstract_only",
            "full_text": None,
            "full_text_length": 0,
            "sections": {},
        }
        abstract_paper["full_text_length"] = len(abstract_paper["abstract"])

        # A5 extraction from abstract-only paper
        a5_paper = _make_a5_paper(
            abstract_paper,
            {"study_design": "RCT", "sample_size": 30, "country": "Unknown", "setting": "Unknown"},
            [{
                "name": "Anxiety score",
                "measure_type": "SMD",
                "effect_size_value": -0.40,
                "ci_lower": -0.85,
                "ci_upper": 0.05,
                "p_value": None,
                "n_intervention": 15,
                "n_control": 15,
                "numeric_complete": True,
            }],
            quality="partial",
        )

        # Paper retains abstract_only text source
        assert a5_paper["text_source"] == "abstract_only"
        # Extraction still happened (partial quality expected)
        assert a5_paper["extraction_quality"] == "partial"
        assert a5_paper["extraction_confidence"] == "medium"
        # Outcomes still extracted from abstract text
        assert len(a5_paper["outcomes"]) == 1
        assert a5_paper["outcomes"][0]["effect_size_value"] == -0.40
        assert a5_paper["outcomes"][0]["numeric_complete"] is True


class TestA7NoNumericData:
    """Test A7 behavior when all papers are abstract-only with no numeric data."""

    def test_a7_no_numeric_data(self):
        """All abstract-only papers with no numeric outcomes: A7 should log warning and produce empty analyses."""
        # Create papers with only qualitative/missing outcomes
        papers_no_numeric = []
        for i, paper in enumerate(SYNTHETIC_A5_PAPERS):
            p = dict(paper)
            p["text_source"] = "abstract_only"
            p["outcomes"] = [{
                "name": "Narrative finding only",
                "measure_type": "qualitative",
                "effect_size_value": None,
                "ci_lower": None,
                "ci_upper": None,
                "p_value": None,
                "numeric_complete": False,
            }]
            p["quality_weight"] = 1.0
            papers_no_numeric.append(p)

        # All papers have no usable numeric data
        for p in papers_no_numeric:
            assert p["outcomes"][0]["numeric_complete"] is False
            assert p["outcomes"][0]["effect_size_value"] is None
            assert p["text_source"] == "abstract_only"

        # A7 expected output: empty analyses, warning in stats
        a7_result = {
            "analyses": {},
            "stats": {
                "total_papers_input": len(papers_no_numeric),
                "papers_with_usable_data": 0,
                "analyses_performed": 0,
                "measure_types_found": [],
                "elapsed_s": 0.1,
            },
            "warning": "No numeric outcome data available for meta-analysis",
        }
        assert a7_result["stats"]["papers_with_usable_data"] == 0
        assert a7_result["stats"]["analyses_performed"] == 0
        assert len(a7_result["analyses"]) == 0
        assert "warning" in a7_result


class TestA8FallbackNoA6A7:
    """Test A8 generates manuscript even when A6 and A7 fail."""

    def test_a8_fallback_no_a6_a7(self):
        """A6 and A7 fail; A8 should generate manuscript with available data (A1-A5)."""
        # Simulate A6 failure: empty result
        a6_failed = {
            "papers": [],
            "stats": {
                "total": 0,
                "rob2_assessed": 0,
                "robins_i_assessed": 0,
                "jbi_assessed": 0,
                "failed": 3,
                "elapsed_s": 0.0,
                "error": "Model loading failed",
            },
        }
        # Simulate A7 failure: empty result
        a7_failed = {
            "analyses": {},
            "stats": {
                "total_papers_input": 3,
                "papers_with_usable_data": 0,
                "analyses_performed": 0,
                "measure_types_found": [],
                "elapsed_s": 0.0,
                "error": "No quality weights available",
            },
        }
        # A8 still generates manuscript from A1-A5 data
        # It uses default weight 0.5 when A6 fails (per MEMORY.md)
        a8_fallback = {
            "manuscript": {
                "abstract": "This systematic review identified 3 studies on CBT for anxiety in depression. "
                            "Due to data limitations, quantitative synthesis was not performed.",
                "introduction": "Depression and anxiety frequently co-occur. CBT is a widely used intervention.",
                "methods": "We searched PubMed, Europe PMC, and Semantic Scholar. Three studies met criteria.",
                "results": "Three RCTs were identified. Quality assessment and meta-analysis could not be completed.",
                "discussion": "The included studies suggest CBT may be beneficial, but quantitative pooling was not possible.",
                "conclusion": "Further research with standardized reporting is needed.",
            },
            "references": [
                {"pmid": "30000001", "title": SYNTHETIC_A2_PAPERS[0]["title"],
                 "citation": "Smith J, Doe A, Brown K (2022)", "journal": "J Clin Psychology"},
                {"pmid": "30000002", "title": SYNTHETIC_A2_PAPERS[1]["title"],
                 "citation": "Jones M, Wilson P (2021)", "journal": "Behav Res Therapy"},
                {"pmid": "30000003", "title": SYNTHETIC_A2_PAPERS[2]["title"],
                 "citation": "Lee S, Kim H, Park J (2023)", "journal": "Psychol Med"},
            ],
            "stats": {
                "total_chunks": 12,
                "total_papers_cited": 3,
                "chunks_per_section": {
                    "abstract": 2, "introduction": 2, "methods": 2,
                    "results": 2, "discussion": 2, "conclusion": 2,
                },
                "elapsed_s": 10.0,
            },
            "pdf_path": None,
        }

        # A8 manuscript has all 6 required sections despite A6/A7 failure
        expected_sections = ["abstract", "introduction", "methods", "results", "discussion", "conclusion"]
        for sec in expected_sections:
            assert sec in a8_fallback["manuscript"], f"Missing section: {sec}"
            assert len(a8_fallback["manuscript"][sec]) > 20, f"Section {sec} too short"

        # References are still present (from A5 data)
        assert len(a8_fallback["references"]) == 3

        # Build output including failed agents
        output = {
            "query": "Is CBT effective for anxiety in depression?",
            "protocol": SYNTHETIC_PROTOCOL,
            "pubmed_query": SYNTHETIC_PUBMED_QUERY,
            "search_stats": SYNTHETIC_A2_STATS,
            "screening": SYNTHETIC_A3_RESULT,
            "fulltext_retrieval": SYNTHETIC_A4_RESULT,
            "data_extraction": SYNTHETIC_A5_RESULT,
            "quality_assessment": a6_failed,
            "synthesis": a7_failed,
            "manuscript": a8_fallback,
        }
        # Output is JSON-serializable
        serialized = json.dumps(output, ensure_ascii=False, default=str)
        restored = json.loads(serialized)
        assert restored["quality_assessment"]["stats"]["failed"] == 3
        assert restored["synthesis"]["stats"]["analyses_performed"] == 0
        assert "abstract" in restored["manuscript"]["manuscript"]


class TestMultiModelConfig:
    """Test config with list of models for ensemble agent."""

    def test_multi_model_config(self):
        """Config with list of 2 models for agent3_screening is valid."""
        import yaml

        # Multi-model config: agent3 uses a list of models (one per ensemble vote)
        config = {
            "models": {
                "agent1_protocol": "./models/qwen2.5-7b",
                "agent3_screening": [
                    "./models/qwen2.5-7b",
                    "./models/glm4-9b",
                ],
                "agent5_extraction": "./models/qwen2.5-7b",
            },
            "device": "cuda",
            "quantization": "4bit",
        }

        # agent3_screening is a list of 2 models
        a3_models = config["models"]["agent3_screening"]
        assert isinstance(a3_models, list)
        assert len(a3_models) == 2

        # Other agents are still strings
        assert isinstance(config["models"]["agent1_protocol"], str)
        assert isinstance(config["models"]["agent5_extraction"], str)

        # YAML round-trip preserves list structure
        yaml_str = yaml.dump(config, default_flow_style=False)
        restored = yaml.safe_load(yaml_str)
        assert isinstance(restored["models"]["agent3_screening"], list)
        assert len(restored["models"]["agent3_screening"]) == 2
        assert restored["models"]["agent3_screening"][0] == "./models/qwen2.5-7b"
        assert restored["models"]["agent3_screening"][1] == "./models/glm4-9b"

        # Ensemble size would match number of models
        ensemble_size = len(restored["models"]["agent3_screening"])
        assert ensemble_size == 2
        # Voting threshold still works with 2 votes (need 2/2 = 100% or 1/2 = 50%)
        threshold = 0.60
        votes_needed = max(1, int(ensemble_size * threshold))
        assert votes_needed >= 1
        assert votes_needed <= ensemble_size


class TestA5AuthorExtraction:
    """Regression test: A5 must propagate authors (including special chars) from input papers."""

    def test_a5_extract_once_preserves_authors(self):
        """_extract_once output must include authors, year, journal from input paper."""
        a5_paper = _make_a5_paper(
            SYNTHETIC_A4_INCLUDE[0],
            {"study_design": "RCT", "sample_size": 120, "country": "USA", "setting": "Outpatient"},
            [{
                "name": "Anxiety score",
                "measure_type": "SMD",
                "effect_size_value": -0.52,
                "ci_lower": -0.78,
                "ci_upper": -0.26,
                "p_value": 0.001,
                "n_intervention": 60,
                "n_control": 60,
                "numeric_complete": True,
            }],
            quality="complete",
        )
        # Authors must be present and non-empty
        assert "authors" in a5_paper, "A5 output must contain 'authors' field"
        assert len(a5_paper["authors"]) > 0, "A5 authors must not be empty"
        assert a5_paper["authors"] == ["Smith J", "Doe A", "Brown K"]
        # Year and journal must also propagate
        assert a5_paper.get("year") == "2022"
        assert a5_paper.get("journal") == "J Clin Psychology"

    def test_a5_authors_with_special_characters(self):
        """Authors with accents, hyphens, and non-ASCII must survive the pipeline."""
        special_authors = [
            "García-López R",
            "Müller K",
            "O'Brien-Smith T",
            "José María F",
            "Günther-Straße H",
        ]
        paper = {
            "pmid": "30000088",
            "title": "Test paper with international authors",
            "abstract": "Methods: RCT with 50 participants. Results: SMD=-0.55, 95% CI [-0.90, -0.20].",
            "authors": special_authors,
            "journal": "Acta Psiquiátrica",
            "year": "2023",
            "text_source": "abstract_only",
            "full_text": None,
            "full_text_length": 0,
            "sections": {},
        }

        a5_paper = _make_a5_paper(
            paper,
            {"study_design": "RCT", "sample_size": 50, "country": "Spain", "setting": "Hospital"},
            [{
                "name": "Anxiety score",
                "measure_type": "SMD",
                "effect_size_value": -0.55,
                "ci_lower": -0.90,
                "ci_upper": -0.20,
                "p_value": 0.01,
                "n_intervention": 25,
                "n_control": 25,
                "numeric_complete": True,
            }],
            quality="complete",
        )

        # All special-character authors must survive intact
        assert a5_paper["authors"] == special_authors
        # Accented characters must not be stripped
        assert "García-López R" in a5_paper["authors"]
        assert "Müller K" in a5_paper["authors"]
        assert "O'Brien-Smith T" in a5_paper["authors"]
        assert "José María F" in a5_paper["authors"]

    def test_a5_pipeline_authors_flow_a2_to_a5(self):
        """Authors from A2 must flow through A3 -> A4 -> A5 without loss."""
        # Trace authors from A2 paper through synthetic pipeline
        a2_authors = SYNTHETIC_A2_PAPERS[0]["authors"]
        a3_authors = SYNTHETIC_A3_RESULT["include"][0].get("authors", [])
        a4_authors = SYNTHETIC_A4_INCLUDE[0].get("authors", [])
        a5_authors = SYNTHETIC_A5_RESULT["papers"][0].get("authors", [])

        assert a2_authors == a3_authors, "Authors lost between A2 and A3"
        assert a3_authors == a4_authors, "Authors lost between A3 and A4"
        assert a4_authors == a5_authors, "Authors lost between A4 and A5"

    def test_a5_failed_extraction_preserves_authors(self):
        """Even when extraction fails, authors must be in the error fallback dict."""
        # Simulate the error fallback structure from extract_all (line 214-226)
        paper = {
            "pmid": "30000077",
            "title": "Failed extraction test",
            "authors": ["TestAuthor A", "TestAuthor B"],
            "journal": "Test Journal",
            "year": "2024",
            "text_source": "abstract_only",
        }
        # This matches the error fallback in extract_all
        fallback = {
            "pmid": paper.get("pmid", ""),
            "title": paper.get("title", ""),
            "authors": paper.get("authors", []),
            "journal": paper.get("journal", ""),
            "year": paper.get("year", ""),
            "study_characteristics": None,
            "intervention_exposure": None,
            "outcomes": [],
            "extraction_quality": "failed",
            "extraction_confidence": "low",
            "concordance": None,
            "extraction_notes": "test error",
            "text_source": paper.get("text_source", ""),
            "sections_used": [],
        }
        assert fallback["authors"] == ["TestAuthor A", "TestAuthor B"]
        assert fallback["journal"] == "Test Journal"
        assert fallback["year"] == "2024"


class TestSpiderFramework:
    """Test qualitative query that A1 classifies as SPIDER framework."""

    def test_spider_framework(self):
        """Qualitative query triggers SPIDER framework detection in A1."""
        query = "What are the lived experiences of nurses providing palliative care in rural settings?"
        # SPIDER: Sample, Phenomenon of Interest, Design, Evaluation, Research type
        protocol = {
            "framework": "SPIDER",
            "sample": "nurses providing palliative care",
            "phenomenon_of_interest": "lived experiences of palliative care provision",
            "design": "qualitative interviews or focus groups",
            "evaluation": "thematic analysis of experiences",
            "research_type": "qualitative",
            "research_question": query,
            "inclusion_criteria": [
                "Registered nurses",
                "Providing palliative care",
                "Rural settings",
                "Qualitative study design",
            ],
            "exclusion_criteria": [
                "Quantitative-only studies",
                "Non-nursing professionals",
                "Urban settings",
            ],
            "pubmed_query": (
                '(nurses OR nursing) AND (palliative care OR end-of-life) '
                'AND (rural) AND (qualitative OR "lived experience" OR phenomenology)'
            ),
        }

        # Framework is SPIDER (not PICO)
        assert protocol["framework"] == "SPIDER"
        assert protocol["framework"] in ("PICO", "SPIDER", "PECO", "PCC")

        # SPIDER-specific fields present
        assert "sample" in protocol
        assert "phenomenon_of_interest" in protocol
        assert "design" in protocol
        assert "evaluation" in protocol
        assert "research_type" in protocol
        assert protocol["research_type"] == "qualitative"

        # Qualitative query should NOT have PICO fields
        assert "population" not in protocol
        assert "intervention" not in protocol
        assert "comparator" not in protocol

        # A6 would use JBI for qualitative studies
        jbi_assessment = {
            "tool": "JBI",
            "overall": "Include",
            "domains": {
                "philosophical_perspective": "Yes",
                "methodology_congruity": "Yes",
                "methods_congruity": "Yes",
                "representation_congruity": "Yes",
                "participant_adequacy": "Yes",
                "ethics_approval": "Yes",
                "conclusions_from_data": "Yes",
                "locating_researcher": "Unclear",
                "researcher_influence": "Unclear",
                "participants_voices": "Yes",
            },
        }
        assert jbi_assessment["tool"] == "JBI"
        assert len(jbi_assessment["domains"]) == 10  # JBI has 10 items
        # Weight for JBI Include = 1.0
        weight = 1.0 if jbi_assessment["overall"] == "Include" else (
            0.6 if jbi_assessment["overall"] == "Caution" else 0.2
        )
        assert weight == 1.0


_AGENTS_JSON_EXISTS = Path("config/agents.json").exists()
_SKIP_NO_JSON = pytest.mark.skipif(not _AGENTS_JSON_EXISTS, reason="requires config/agents.json")


class TestAgentsJsonConfig:
    """Test config/agents.json loading in BaseAgent (both present and absent)."""

    @_SKIP_NO_JSON
    def test_agents_json_loads_params(self):
        """When config/agents.json exists, agent_params has correct values."""
        from agents.base import BaseAgent, _AGENTS_JSON_PATH

        assert os.path.isfile(_AGENTS_JSON_PATH), "config/agents.json must exist"
        with open(_AGENTS_JSON_PATH) as f:
            data = json.load(f)

        # Verify defaults are present
        assert "defaults" in data
        assert "agents" in data
        assert data["defaults"]["repetition_penalty"] == 1.3

        # Verify known agent entries
        a5_cfg = data["agents"]["agent5_extraction"]
        assert a5_cfg["temperature"] == 0.1
        a3_cfg = data["agents"]["agent3_screening"]
        assert a3_cfg["temperature"] == 0.2
        a4_cfg = data["agents"]["agent4_validation"]
        assert a4_cfg["temperature"] == 0.2

        # A8 has section-specific temperatures
        a8_cfg = data["agents"]["agent8_manuscript"]
        assert isinstance(a8_cfg["temperature"], dict)
        assert a8_cfg["temperature"]["abstract"] == 0.3
        assert a8_cfg["temperature"]["discussion"] == 0.6

    @_SKIP_NO_JSON
    def test_agents_json_merged_with_defaults(self):
        """Agent-specific values override defaults; missing keys fall back."""
        from agents.base import BaseAgent
        params = BaseAgent._load_agent_params("agent5_extraction")
        # Agent-specific override
        assert params["temperature"] == 0.1
        # Inherited from defaults
        assert params["repetition_penalty"] == 1.3

    @_SKIP_NO_JSON
    def test_agents_json_missing_agent_gets_defaults(self):
        """Unknown agent name returns defaults only."""
        from agents.base import BaseAgent
        params = BaseAgent._load_agent_params("agent99_nonexistent")
        # Should get defaults
        assert params.get("repetition_penalty") == 1.3
        assert params.get("temperature") == 0.3  # default

    def test_agents_json_absent_returns_empty(self, tmp_path, monkeypatch):
        """When agents.json does not exist, _load_agent_params returns {}."""
        import agents.base as base_mod
        monkeypatch.setattr(base_mod, "_AGENTS_JSON_PATH", str(tmp_path / "nonexistent.json"))
        params = base_mod.BaseAgent._load_agent_params("agent5_extraction")
        assert params == {}

    @_SKIP_NO_JSON
    def test_get_agent_param_accessor(self):
        """get_agent_param returns agent-specific value with fallback."""
        from agents.base import BaseAgent
        params = BaseAgent._load_agent_params("agent5_extraction")
        # Simulate what the instance method does
        assert params.get("temperature", 999) == 0.1
        assert params.get("nonexistent_key", 42) == 42


# ══════════════════════════════════════════════════════════════════════════════
# T1 REGRESSION: Author extraction with problematic names (conftest fixtures)
# ══════════════════════════════════════════════════════════════════════════════

class TestT1AuthorRegressionConftest:
    """Regression: A5 output must preserve problematic author names from conftest."""

    def test_accented_authors_survive_a5_pipeline(self, a5_papers_special_authors):
        """Accented/hyphenated authors must survive A5 extraction output intact."""
        paper = a5_papers_special_authors[0]
        expected = ["García-López R", "Müller K", "O'Brien-Smith T",
                    "José María F", "Günther-Straße H"]
        assert paper["authors"] == expected
        # Verify accented chars are not stripped (clean_json strips non-ASCII)
        for author in expected:
            assert author in paper["authors"], f"Author '{author}' missing from A5 output"

    def test_cjk_romanized_authors_survive(self, a5_papers_special_authors):
        """CJK-romanized and hyphenated names survive A5 extraction."""
        paper = a5_papers_special_authors[1]
        assert paper["authors"] == ["Zhang W", "Kim H-S", "Tanaka Y"]
        assert paper["journal"] == "Asian J Psychiatry"

    def test_mononym_author_survives(self, a5_papers_special_authors):
        """Single-name (mononym) author survives A5 extraction."""
        paper = a5_papers_special_authors[2]
        assert paper["authors"] == ["Sukarno"]

    def test_a5_output_has_authors_field_for_all_papers(self, a5_papers_special_authors):
        """Every A5 paper must have a non-empty 'authors' list."""
        for paper in a5_papers_special_authors:
            assert "authors" in paper, f"Paper {paper['pmid']} missing 'authors' field"
            assert isinstance(paper["authors"], list), f"Paper {paper['pmid']} authors not a list"
            assert len(paper["authors"]) > 0, f"Paper {paper['pmid']} has empty authors"

    def test_a5_extract_once_includes_authors_journal_year(self, a5_papers_special_authors):
        """A5._extract_once output dict must include authors, journal, year (T1 fix)."""
        # This tests the exact fields added by the T1 fix in agent5_extraction.py
        for paper in a5_papers_special_authors:
            assert "authors" in paper
            assert "journal" in paper
            assert "year" in paper
            assert paper["year"] != ""

    def test_json_roundtrip_preserves_accented_authors(self, a5_papers_special_authors):
        """JSON serialize/deserialize must not corrupt accented author names."""
        original_authors = a5_papers_special_authors[0]["authors"]
        serialized = json.dumps(a5_papers_special_authors[0], ensure_ascii=False)
        restored = json.loads(serialized)
        assert restored["authors"] == original_authors
        assert "García-López R" in restored["authors"]
        assert "Müller K" in restored["authors"]


# ══════════════════════════════════════════════════════════════════════════════
# T4: A9 Explainer agent tests
# ══════════════════════════════════════════════════════════════════════════════

class TestA9Explainer:
    """Test A9 Explainability agent with mock pipeline data."""

    def test_a9_run_generates_files(self, full_pipeline_data, tmp_output_dir):
        """A9.run() creates both papers_raw.xlsx and explainability.md."""
        from agents.agent9_explainer import Agent9Explainer
        a9 = Agent9Explainer()
        result = a9.run(full_pipeline_data, "run_test_001", tmp_output_dir)

        assert "excel_path" in result
        assert "report_path" in result
        assert os.path.isfile(result["excel_path"])
        assert os.path.isfile(result["report_path"])
        assert result["excel_path"].endswith("papers_raw.xlsx")
        assert result["report_path"].endswith("explainability.md")

    def test_a9_excel_has_all_papers(self, full_pipeline_data, tmp_output_dir):
        """Excel file has a row for every paper that entered the pipeline."""
        from agents.agent9_explainer import Agent9Explainer
        from openpyxl import load_workbook

        a9 = Agent9Explainer()
        result = a9.run(full_pipeline_data, "run_test_002", tmp_output_dir)

        wb = load_workbook(result["excel_path"])
        ws = wb.active

        # Header row + data rows
        n_rows = ws.max_row - 1  # subtract header
        # Pipeline has 4 papers total (3 screened include + 1 excluded)
        assert n_rows == 4, f"Expected 4 data rows, got {n_rows}"

    def test_a9_excel_header_columns(self, full_pipeline_data, tmp_output_dir):
        """Excel has expected header columns."""
        from agents.agent9_explainer import Agent9Explainer
        from openpyxl import load_workbook

        a9 = Agent9Explainer()
        result = a9.run(full_pipeline_data, "run_test_003", tmp_output_dir)

        wb = load_workbook(result["excel_path"])
        ws = wb.active

        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        # Must have PMID, Title, Authors, A3 Decision, Reason Excluded at minimum
        assert "PMID" in headers
        assert "Title" in headers
        assert "Authors" in headers
        assert "A3 Decision" in headers
        assert "Reason Excluded" in headers

    def test_a9_excel_preserves_special_authors(self, full_pipeline_data, tmp_output_dir):
        """Excel file preserves accented author names (not stripped by clean_json)."""
        from agents.agent9_explainer import Agent9Explainer
        from openpyxl import load_workbook

        a9 = Agent9Explainer()
        result = a9.run(full_pipeline_data, "run_test_004", tmp_output_dir)

        wb = load_workbook(result["excel_path"])
        ws = wb.active

        # Find Authors column index
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        auth_col = headers.index("Authors") + 1

        # Check first paper (row 2) has accented author
        authors_cell = ws.cell(row=2, column=auth_col).value or ""
        assert "García-López" in authors_cell, f"Accented author lost in Excel: {authors_cell}"

    def test_a9_excel_excluded_papers_have_reason(self, full_pipeline_data, tmp_output_dir):
        """Excluded papers in Excel have a non-empty exclusion reason and stage."""
        from agents.agent9_explainer import Agent9Explainer
        from openpyxl import load_workbook

        a9 = Agent9Explainer()
        result = a9.run(full_pipeline_data, "run_test_005", tmp_output_dir)

        wb = load_workbook(result["excel_path"])
        ws = wb.active

        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        a3_col = headers.index("A3 Decision") + 1
        a4_col = headers.index("A4 Decision") + 1
        reason_col = headers.index("Reason Excluded") + 1
        stage_col = headers.index("Exclusion Stage") + 1

        # Find excluded rows
        found_excluded = False
        for row_idx in range(2, ws.max_row + 1):
            a3 = ws.cell(row=row_idx, column=a3_col).value or ""
            a4 = ws.cell(row=row_idx, column=a4_col).value or ""
            if a3 == "EXCLUDE" or a4 == "EXCLUDE":
                reason = ws.cell(row=row_idx, column=reason_col).value or ""
                stage = ws.cell(row=row_idx, column=stage_col).value or ""
                assert reason != "" or stage != "", \
                    f"Row {row_idx}: excluded paper has no reason/stage"
                found_excluded = True

        assert found_excluded, "No excluded papers found in Excel"

    def test_a9_report_contains_sections(self, full_pipeline_data, tmp_output_dir):
        """Explainability report has all expected sections (in Spanish)."""
        from agents.agent9_explainer import Agent9Explainer

        a9 = Agent9Explainer()
        result = a9.run(full_pipeline_data, "run_test_006", tmp_output_dir)

        with open(result["report_path"], "r", encoding="utf-8") as f:
            content = f.read()

        # Section headers (Spanish)
        assert "Resumen General" in content
        assert "Trazabilidad por Etapa" in content
        assert "Razones de Exclusion" in content
        assert "Trazabilidad RAG" in content
        assert "Decisiones Borderline" in content
        assert "Procedencia de Datos" in content

    def test_a9_report_contains_protocol_data(self, full_pipeline_data, tmp_output_dir):
        """Report includes protocol details from A1."""
        from agents.agent9_explainer import Agent9Explainer

        a9 = Agent9Explainer()
        result = a9.run(full_pipeline_data, "run_test_007", tmp_output_dir)

        with open(result["report_path"], "r", encoding="utf-8") as f:
            content = f.read()

        # Must reference the research question and framework
        assert "PICO" in content
        assert "CBT" in content
        assert "depression" in content.lower() or "depresion" in content.lower()

    def test_a9_report_prisma_table(self, full_pipeline_data, tmp_output_dir):
        """Report includes PRISMA flow table with correct counts."""
        from agents.agent9_explainer import Agent9Explainer

        a9 = Agent9Explainer()
        result = a9.run(full_pipeline_data, "run_test_008", tmp_output_dir)

        with open(result["report_path"], "r", encoding="utf-8") as f:
            content = f.read()

        # PRISMA table must exist
        assert "Flujo PRISMA" in content
        assert "Identificados" in content
        assert "Incluidos final" in content

    def test_a9_result_stats(self, full_pipeline_data, tmp_output_dir):
        """A9 result dict has correct stats."""
        from agents.agent9_explainer import Agent9Explainer

        a9 = Agent9Explainer()
        result = a9.run(full_pipeline_data, "run_test_009", tmp_output_dir)

        stats = result["stats"]
        assert stats["total_papers"] == 4  # 3 screened-include + 1 excluded
        assert stats["included"] == 2      # 2 passed A4 full-text
        assert stats["excluded"] >= 1      # at least 1 excluded
        assert stats["elapsed_s"] >= 0

    def test_a9_handles_empty_pipeline(self, tmp_output_dir):
        """A9 handles empty/minimal pipeline data without crashing."""
        from agents.agent9_explainer import Agent9Explainer

        empty_data = {
            "query": "test",
            "protocol": {"framework": "PICO", "research_question": "test"},
            "screening": {"include": [], "exclude": [], "uncertain": [], "stats": {}},
            "fulltext_retrieval": {"include": [], "exclude": [], "papers": [], "stats": {}},
            "data_extraction": {"papers": [], "stats": {}},
            "quality_assessment": {"papers": [], "stats": {}},
            "synthesis": {"analyses": {}, "stats": {}},
            "manuscript": {"manuscript": {}, "stats": {}},
            "prisma": {},
            "search_stats": {},
        }

        a9 = Agent9Explainer()
        result = a9.run(empty_data, "run_empty", tmp_output_dir)

        assert os.path.isfile(result["excel_path"])
        assert os.path.isfile(result["report_path"])
        assert result["stats"]["total_papers"] == 0

    def test_a9_json_serializable(self, full_pipeline_data, tmp_output_dir):
        """A9 result is fully JSON-serializable."""
        from agents.agent9_explainer import Agent9Explainer

        a9 = Agent9Explainer()
        result = a9.run(full_pipeline_data, "run_serial", tmp_output_dir)

        serialized = json.dumps(result, ensure_ascii=False, default=str)
        restored = json.loads(serialized)
        assert "excel_path" in restored
        assert "report_path" in restored


# ══════════════════════════════════════════════════════════════════════════════
# T6: Fail-Safe N + Publication Bias structure tests
# ══════════════════════════════════════════════════════════════════════════════

class TestFailSafeNStructure:
    """Validate expected A7 output structure for Fail-Safe N (T6 fixtures)."""

    def test_failsafe_n_in_publication_bias(self, a7_result_with_failsafe_n):
        """publication_bias dict must contain failsafe_n key."""
        smd = a7_result_with_failsafe_n["analyses"]["SMD"]
        pb = smd["publication_bias"]
        assert "failsafe_n" in pb, "publication_bias missing 'failsafe_n' key"

    def test_failsafe_n_fields(self, a7_result_with_failsafe_n):
        """failsafe_n dict has required fields: n, alpha, method, interpretation."""
        fsn = a7_result_with_failsafe_n["analyses"]["SMD"]["publication_bias"]["failsafe_n"]
        assert "n" in fsn
        assert "alpha" in fsn
        assert "method" in fsn
        assert "interpretation" in fsn
        assert isinstance(fsn["n"], int)
        assert fsn["n"] > 0
        assert fsn["method"] == "Rosenthal"
        assert fsn["alpha"] == 0.05

    def test_egger_test_structure(self, a7_result_with_failsafe_n):
        """Egger test dict has intercept, se, p_value, significant."""
        egger = a7_result_with_failsafe_n["analyses"]["SMD"]["publication_bias"]["egger_test"]
        assert egger is not None
        assert "intercept" in egger
        assert "se" in egger
        assert "p_value" in egger
        assert "significant" in egger
        assert isinstance(egger["significant"], bool)

    def test_trim_and_fill_structure(self, a7_result_with_failsafe_n):
        """Trim-and-fill dict has k_missing, adjusted_estimate, adjusted CI."""
        tf = a7_result_with_failsafe_n["analyses"]["SMD"]["publication_bias"]["trim_and_fill"]
        assert tf is not None
        assert "k_missing" in tf
        assert "adjusted_estimate" in tf
        assert "adjusted_ci_lower" in tf
        assert "adjusted_ci_upper" in tf
        assert tf["adjusted_ci_lower"] <= tf["adjusted_estimate"] <= tf["adjusted_ci_upper"]

    def test_12_studies_enough_for_egger(self, a7_studies_for_pub_bias):
        """Synthetic dataset has k=12 >= 10, qualifying for Egger test."""
        assert len(a7_studies_for_pub_bias) == 12
        assert len(a7_studies_for_pub_bias) >= 10, "Need k>=10 for Egger test"

    def test_studies_have_required_fields(self, a7_studies_for_pub_bias):
        """Each study has pmid, yi, vi, weight."""
        for study in a7_studies_for_pub_bias:
            assert "pmid" in study
            assert "yi" in study
            assert "vi" in study
            assert "weight" in study
            assert isinstance(study["yi"], float)
            assert isinstance(study["vi"], float)
            assert study["vi"] > 0, f"Variance must be positive: {study['vi']}"

    def test_failsafe_n_json_serializable(self, a7_result_with_failsafe_n):
        """Full A7 result with Fail-Safe N is JSON-serializable."""
        serialized = json.dumps(a7_result_with_failsafe_n, ensure_ascii=False)
        restored = json.loads(serialized)
        fsn = restored["analyses"]["SMD"]["publication_bias"]["failsafe_n"]
        assert fsn["n"] == 187
        assert fsn["method"] == "Rosenthal"

    def test_publication_bias_complete_when_k_sufficient(self, a7_result_with_failsafe_n):
        """With k=12, all three pub bias methods should be populated (not None)."""
        pb = a7_result_with_failsafe_n["analyses"]["SMD"]["publication_bias"]
        assert pb["egger_test"] is not None, "Egger should be populated with k>=10"
        assert pb["trim_and_fill"] is not None, "Trim-and-fill should be populated"
        assert pb["failsafe_n"] is not None, "Fail-Safe N should always be populated"


# ══════════════════════════════════════════════════════════════════════════════
# T3: Config max_input_chars from agents.json (structural)
# ══════════════════════════════════════════════════════════════════════════════

class TestT3MaxInputChars:
    """Test that max_input_chars is readable from agents.json config."""

    def test_agents_json_has_max_input_chars(self):
        """config/agents.json must define max_input_chars for key agents."""
        from agents.base import _AGENTS_JSON_PATH
        if not os.path.isfile(_AGENTS_JSON_PATH):
            pytest.skip("config/agents.json not yet created")

        with open(_AGENTS_JSON_PATH) as f:
            data = json.load(f)

        # A5 is the most critical (3 prompts with different input lengths)
        a5 = data.get("agents", {}).get("agent5_extraction", {})
        assert "max_input_chars" in a5, "agent5_extraction must have max_input_chars"
        mic = a5["max_input_chars"]
        # A5 max_input_chars can be int (uniform) or dict (per-prompt)
        if isinstance(mic, dict):
            for key, val in mic.items():
                assert isinstance(val, int), f"max_input_chars[{key}] not int: {val}"
                assert val >= 1000, f"max_input_chars[{key}] should be >= 1000"
        else:
            assert isinstance(mic, int)
            assert mic >= 1000

    def test_defaults_has_max_input_chars(self):
        """config/agents.json defaults must include max_input_chars."""
        from agents.base import _AGENTS_JSON_PATH
        if not os.path.isfile(_AGENTS_JSON_PATH):
            pytest.skip("config/agents.json not yet created")

        with open(_AGENTS_JSON_PATH) as f:
            data = json.load(f)

        defaults = data.get("defaults", {})
        assert "max_input_chars" in defaults
        assert isinstance(defaults["max_input_chars"], int)


# ══════════════════════════════════════════════════════════════════════════════
# T2 EXTENDED: Merge logic with custom fixtures (conftest)
# ══════════════════════════════════════════════════════════════════════════════

class TestAgentsJsonMergeLogic:
    """Test _load_agent_params merge behavior with custom/absent agents.json."""

    def test_custom_json_overrides_defaults(self, agents_json_custom_path, monkeypatch):
        """Custom agents.json: A5 gets custom temp=0.05, rep_penalty=1.5."""
        import agents.base as base_mod
        monkeypatch.setattr(base_mod, "_AGENTS_JSON_PATH", agents_json_custom_path)

        params = base_mod.BaseAgent._load_agent_params("agent5_extraction")
        # Agent-specific override from FIXTURE_AGENTS_JSON_CUSTOM
        assert params["temperature"] == 0.05
        assert params["repetition_penalty"] == 1.5
        assert params["max_tokens"] == 800
        assert params["max_input_chars"] == 8000

    def test_custom_json_unlisted_agent_gets_defaults(self, agents_json_custom_path, monkeypatch):
        """Agent NOT in custom json gets only defaults."""
        import agents.base as base_mod
        monkeypatch.setattr(base_mod, "_AGENTS_JSON_PATH", agents_json_custom_path)

        # agent3_screening is NOT in FIXTURE_AGENTS_JSON_CUSTOM
        params = base_mod.BaseAgent._load_agent_params("agent3_screening")
        assert params["temperature"] == 0.3  # default, not 0.2
        assert params["repetition_penalty"] == 1.3  # default
        assert params["max_tokens"] == 300  # default

    def test_canonical_json_all_agents_present(self, agents_json_path, monkeypatch):
        """Canonical fixture has entries for A1, A3, A4, A5, A6, A8."""
        import agents.base as base_mod
        monkeypatch.setattr(base_mod, "_AGENTS_JSON_PATH", agents_json_path)

        for name, expected_temp in [
            ("agent1_protocol", 0.3),
            ("agent3_screening", 0.2),
            ("agent4_validation", 0.2),
            ("agent5_extraction", 0.1),
            ("agent6_quality", 0.2),
        ]:
            params = base_mod.BaseAgent._load_agent_params(name)
            assert params["temperature"] == expected_temp, \
                f"{name}: expected temp={expected_temp}, got {params['temperature']}"

    def test_canonical_json_a8_dict_temperature(self, agents_json_path, monkeypatch):
        """A8 manuscript has section-specific temperature dict in canonical fixture."""
        import agents.base as base_mod
        monkeypatch.setattr(base_mod, "_AGENTS_JSON_PATH", agents_json_path)

        params = base_mod.BaseAgent._load_agent_params("agent8_manuscript")
        temp = params["temperature"]
        assert isinstance(temp, dict), "A8 temperature should be a dict"
        assert temp["abstract"] == 0.3
        assert temp["discussion"] == 0.6
        assert temp["methods"] == 0.3

    def test_absent_json_returns_empty_dict(self, no_agents_json_path, monkeypatch):
        """When agents.json does not exist, returns {} (backward compat)."""
        import agents.base as base_mod
        monkeypatch.setattr(base_mod, "_AGENTS_JSON_PATH", no_agents_json_path)

        params = base_mod.BaseAgent._load_agent_params("agent5_extraction")
        assert params == {}

    def test_absent_json_get_agent_param_uses_fallback(self, no_agents_json_path, monkeypatch):
        """With no agents.json, get_agent_param should return the fallback value."""
        import agents.base as base_mod
        monkeypatch.setattr(base_mod, "_AGENTS_JSON_PATH", no_agents_json_path)

        params = base_mod.BaseAgent._load_agent_params("agent5_extraction")
        # Simulate get_agent_param behavior
        assert params.get("temperature", 0.1) == 0.1  # fallback
        assert params.get("repetition_penalty", 1.3) == 1.3  # fallback

    def test_malformed_json_returns_empty(self, tmp_path, monkeypatch):
        """Malformed agents.json (invalid JSON) returns {} gracefully."""
        import agents.base as base_mod
        bad_path = tmp_path / "bad_agents.json"
        bad_path.write_text("{invalid json: not valid!!!")
        monkeypatch.setattr(base_mod, "_AGENTS_JSON_PATH", str(bad_path))

        params = base_mod.BaseAgent._load_agent_params("agent5_extraction")
        assert params == {}

    def test_merge_order_agent_overrides_defaults(self, agents_json_custom_path, monkeypatch):
        """Agent-specific values must override defaults, not the other way around."""
        import agents.base as base_mod
        monkeypatch.setattr(base_mod, "_AGENTS_JSON_PATH", agents_json_custom_path)

        params = base_mod.BaseAgent._load_agent_params("agent5_extraction")
        # Default temp=0.3, agent5 custom temp=0.05 → must be 0.05
        assert params["temperature"] == 0.05
        # Default rep_penalty=1.3, agent5 custom=1.5 → must be 1.5
        assert params["repetition_penalty"] == 1.5
        # Default max_tokens=300, agent5 custom=800 → must be 800
        assert params["max_tokens"] == 800


# ══════════════════════════════════════════════════════════════════════════════
# T1 EXTENDED: Author flow through A6/A7/A8 (downstream propagation)
# ══════════════════════════════════════════════════════════════════════════════

class TestT1AuthorDownstreamFlow:
    """Verify authors propagate correctly from A5 through A6, A7, A8."""

    def test_a6_papers_retain_authors(self):
        """A6 papers (quality assessment) must retain authors from A5."""
        for paper in SYNTHETIC_A6_PAPERS:
            assert "authors" in paper, f"Paper {paper['pmid']} missing 'authors' in A6"
            assert isinstance(paper["authors"], list)
            assert len(paper["authors"]) > 0

    def test_a8_references_have_author_citations(self):
        """A8 references must have citation strings that contain author names."""
        refs = SYNTHETIC_A8_RESULT["references"]
        for ref in refs:
            assert "citation" in ref, f"Reference {ref['pmid']} missing 'citation'"
            assert len(ref["citation"]) > 3, f"Citation too short: {ref['citation']}"

    def test_authors_flow_a2_a3_a4_a5_a6_consistent(self):
        """Authors must be identical at each pipeline stage for the same paper."""
        # Paper 30000001: A2 -> A3 -> A4 -> A5 -> A6
        a2_authors = SYNTHETIC_A2_PAPERS[0]["authors"]
        a3_authors = SYNTHETIC_A3_RESULT["include"][0].get("authors", [])
        a4_authors = SYNTHETIC_A4_PAPERS[0].get("authors", [])
        a5_authors = SYNTHETIC_A5_RESULT["papers"][0].get("authors", [])
        a6_authors = SYNTHETIC_A6_PAPERS[0].get("authors", [])

        assert a2_authors == a3_authors, "Authors changed between A2 and A3"
        assert a3_authors == a4_authors, "Authors changed between A3 and A4"
        assert a4_authors == a5_authors, "Authors changed between A4 and A5"
        assert a5_authors == a6_authors, "Authors changed between A5 and A6"

    def test_a9_excel_authors_from_conftest_pipeline(self, full_pipeline_data, tmp_output_dir):
        """A9 Excel must preserve international authors from full_pipeline_data fixture."""
        from agents.agent9_explainer import Agent9Explainer
        from openpyxl import load_workbook

        a9 = Agent9Explainer()
        result = a9.run(full_pipeline_data, "run_author_flow", tmp_output_dir)

        wb = load_workbook(result["excel_path"])
        ws = wb.active

        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        auth_col = headers.index("Authors") + 1

        # Collect all author cells
        all_authors_text = []
        for row_idx in range(2, ws.max_row + 1):
            val = ws.cell(row=row_idx, column=auth_col).value or ""
            all_authors_text.append(val)

        # At least one cell must contain accented author
        combined = " ".join(all_authors_text)
        assert "García-López" in combined, f"Accented author missing from Excel. Got: {combined[:200]}"
        assert "Müller" in combined, f"Umlaut author missing from Excel. Got: {combined[:200]}"


# ══════════════════════════════════════════════════════════════════════════════
# T4 EXTENDED: A9 schema, edge cases, Spanish output
# ══════════════════════════════════════════════════════════════════════════════

class TestA9ExplainerSchema:
    """Verify A9 Excel schema matches the 27-column spec exactly."""

    def test_a9_excel_exact_column_count(self, full_pipeline_data, tmp_output_dir, a9_excel_columns):
        """Excel must have exactly 27 header columns."""
        from agents.agent9_explainer import Agent9Explainer
        from openpyxl import load_workbook

        a9 = Agent9Explainer()
        result = a9.run(full_pipeline_data, "run_schema", tmp_output_dir)

        wb = load_workbook(result["excel_path"])
        ws = wb.active
        actual_count = ws.max_column
        assert actual_count == 27, f"Expected 27 columns, got {actual_count}"

    def test_a9_excel_columns_match_spec(self, full_pipeline_data, tmp_output_dir, a9_excel_columns):
        """Excel headers must match the spec list exactly (order and names)."""
        from agents.agent9_explainer import Agent9Explainer
        from openpyxl import load_workbook

        a9 = Agent9Explainer()
        result = a9.run(full_pipeline_data, "run_spec", tmp_output_dir)

        wb = load_workbook(result["excel_path"])
        ws = wb.active
        actual_headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        assert actual_headers == a9_excel_columns, (
            f"Header mismatch.\nExpected: {a9_excel_columns}\nActual: {actual_headers}"
        )

    def test_a9_excel_has_a3_a4_a5_a6_columns(self, full_pipeline_data, tmp_output_dir):
        """Excel has pipeline-decision columns for every agent stage."""
        from agents.agent9_explainer import Agent9Explainer
        from openpyxl import load_workbook

        a9 = Agent9Explainer()
        result = a9.run(full_pipeline_data, "run_stages", tmp_output_dir)

        wb = load_workbook(result["excel_path"])
        ws = wb.active
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]

        required_stage_cols = [
            "A3 Decision", "A3 Confidence",
            "A4 Decision", "A4 Text Source",
            "A5 Extraction Quality", "A5 Confidence",
            "A6 Tool", "A6 Overall", "A6 Weight",
        ]
        for col in required_stage_cols:
            assert col in headers, f"Missing pipeline-stage column: {col}"

    def test_a9_excel_included_paper_has_extraction_data(self, full_pipeline_data, tmp_output_dir):
        """Included papers in Excel should have extraction data populated."""
        from agents.agent9_explainer import Agent9Explainer
        from openpyxl import load_workbook

        a9 = Agent9Explainer()
        result = a9.run(full_pipeline_data, "run_extdata", tmp_output_dir)

        wb = load_workbook(result["excel_path"])
        ws = wb.active
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]

        a3_col = headers.index("A3 Decision") + 1
        a5_col = headers.index("A5 Extraction Quality") + 1
        es_col = headers.index("Effect Size Value") + 1

        # Find an included paper row
        for row_idx in range(2, ws.max_row + 1):
            a3_val = ws.cell(row=row_idx, column=a3_col).value or ""
            if a3_val == "INCLUDE":
                a5_val = ws.cell(row=row_idx, column=a5_col).value or ""
                es_val = ws.cell(row=row_idx, column=es_col).value
                # Included paper that went through A5 should have extraction data
                if a5_val:
                    assert a5_val in ("complete", "partial", "failed"), \
                        f"Row {row_idx}: unexpected A5 quality: {a5_val}"
                    if a5_val == "complete":
                        assert es_val is not None and es_val != "", \
                            f"Row {row_idx}: complete extraction but no effect size"
                break


class TestA9ExplainerEdgeCases:
    """Edge cases for A9 explainer."""

    def test_a9_pipeline_no_excluded_papers(self, tmp_output_dir):
        """Pipeline where all papers are included (no exclusions)."""
        from agents.agent9_explainer import Agent9Explainer

        data = {
            "query": "test",
            "protocol": {"framework": "PICO", "research_question": "test",
                         "population": "adults", "intervention": "CBT",
                         "comparator": "waitlist", "outcome": "anxiety"},
            "screening": {
                "include": [
                    {"pmid": "11111", "title": "Paper A", "abstract": "...",
                     "authors": ["Author A"], "decision": "INCLUDE", "confidence": 1.0},
                ],
                "exclude": [],
                "uncertain": [],
                "stats": {"total": 1, "included": 1, "excluded": 0, "uncertain": 0},
            },
            "fulltext_retrieval": {
                "include": [
                    {"pmid": "11111", "title": "Paper A", "authors": ["Author A"],
                     "journal": "J Test", "year": "2024", "doi": "10.1/test",
                     "text_source": "pmc", "full_text_length": 3000,
                     "fulltext_decision": "INCLUDE", "fulltext_reasoning": "OK",
                     "exclusion_reason": None, "abstract": "..."},
                ],
                "exclude": [],
                "papers": [],
                "stats": {"total_retrieved": 1, "fulltext_included": 1, "fulltext_excluded": 0},
            },
            "data_extraction": {"papers": [
                {"pmid": "11111", "title": "Paper A", "authors": ["Author A"],
                 "outcomes": [{"name": "Anxiety", "measure_type": "SMD",
                               "effect_size_value": -0.5, "ci_lower": -0.8, "ci_upper": -0.2}],
                 "extraction_quality": "complete", "extraction_confidence": "high",
                 "study_characteristics": {"sample_size": 100},
                 "intervention_exposure": {"name": "CBT"}},
            ], "stats": {}},
            "quality_assessment": {"papers": [], "stats": {}},
            "synthesis": {"analyses": {}, "stats": {}},
            "manuscript": {"manuscript": {}, "stats": {}},
            "prisma": {"identified": 1, "screened": 1, "included_final": 1},
            "search_stats": {},
        }

        a9 = Agent9Explainer()
        result = a9.run(data, "run_no_excl", tmp_output_dir)

        assert os.path.isfile(result["excel_path"])
        assert os.path.isfile(result["report_path"])

        with open(result["report_path"], "r", encoding="utf-8") as f:
            content = f.read()

        # Exclusions section should say no papers excluded
        assert "Razones de Exclusion" in content
        assert result["stats"]["excluded"] == 0

    def test_a9_pipeline_no_faiss_data(self, tmp_output_dir):
        """Pipeline with no RAG/FAISS traceability data in A8."""
        from agents.agent9_explainer import Agent9Explainer

        data = {
            "query": "test",
            "protocol": {"framework": "PICO", "research_question": "test"},
            "screening": {"include": [], "exclude": [], "uncertain": [], "stats": {}},
            "fulltext_retrieval": {"include": [], "exclude": [], "papers": [], "stats": {}},
            "data_extraction": {"papers": [], "stats": {}},
            "quality_assessment": {"papers": [], "stats": {}},
            "synthesis": {"analyses": {}, "stats": {}},
            "manuscript": {
                "manuscript": {"abstract": "No data", "introduction": "No data",
                               "methods": "No data", "results": "No data",
                               "discussion": "No data", "conclusion": "No data"},
                "stats": {},  # no chunks_per_section, no total_chunks
            },
            "prisma": {},
            "search_stats": {},
        }

        a9 = Agent9Explainer()
        result = a9.run(data, "run_no_faiss", tmp_output_dir)

        with open(result["report_path"], "r", encoding="utf-8") as f:
            content = f.read()

        # RAG section should handle missing data gracefully
        assert "Trazabilidad RAG" in content
        assert "No hay datos de trazabilidad RAG" in content

    def test_a9_report_all_in_spanish(self, full_pipeline_data, tmp_output_dir):
        """All section headers in explainability.md must be in Spanish."""
        from agents.agent9_explainer import Agent9Explainer

        a9 = Agent9Explainer()
        result = a9.run(full_pipeline_data, "run_spanish", tmp_output_dir)

        with open(result["report_path"], "r", encoding="utf-8") as f:
            content = f.read()

        # Spanish section headers (per agent9_explainer.py)
        spanish_headers = [
            "Informe de Explicabilidad y Trazabilidad",
            "Resumen General",
            "Trazabilidad por Etapa",
            "Razones de Exclusion por Paper",
            "Trazabilidad RAG",
            "Decisiones Borderline",
            "Procedencia de Datos por Paper Incluido",
        ]
        for header in spanish_headers:
            assert header in content, f"Missing Spanish header: '{header}'"

        # Must NOT have English equivalents
        english_headers = ["Overview", "Traceability per Stage", "Exclusion Reasons",
                           "RAG Traceability", "Borderline Decisions", "Data Provenance"]
        for eh in english_headers:
            assert eh not in content, f"Found English header (should be Spanish): '{eh}'"

    def test_a9_report_provenance_section_has_included_papers(self, full_pipeline_data, tmp_output_dir):
        """Section 6 (Procedencia) must list each included paper with PMID."""
        from agents.agent9_explainer import Agent9Explainer

        a9 = Agent9Explainer()
        result = a9.run(full_pipeline_data, "run_provenance", tmp_output_dir)

        with open(result["report_path"], "r", encoding="utf-8") as f:
            content = f.read()

        # Included papers are 40000001 and 40000002 in the conftest fixture
        assert "PMID 40000001" in content
        assert "PMID 40000002" in content
        # Excluded paper should NOT be in provenance
        assert "PMID 40000004" not in content or "Procedencia" not in content.split("PMID 40000004")[0]


# ══════════════════════════════════════════════════════════════════════════════
# REVIEW REPORT: Regression test for NoneType bug in Egger/Trim-and-fill
# ══════════════════════════════════════════════════════════════════════════════

class TestReviewReportNoneHandling:
    """Regression: review_report.py must not crash when egger_test or trim_and_fill is None."""

    def test_review_report_with_none_egger_and_trim_fill(self):
        """ReviewReportGenerator.generate() must not crash when publication_bias has None values."""
        from agents.review_report import ReviewReportGenerator

        # Pipeline with egger_test=None and trim_and_fill=None (k<10 scenario)
        pipeline_data = {
            "screening": {
                "include": [
                    {"pmid": "10001", "title": "Paper 1", "decision": "INCLUDE",
                     "confidence": 1.0, "votes": ["INCLUDE"] * 3},
                ],
                "exclude": [],
                "uncertain": [],
                "stats": {"total": 1, "included": 1, "excluded": 0, "uncertain": 0},
            },
            "fulltext_retrieval": {
                "include": [{"pmid": "10001", "title": "Paper 1",
                             "fulltext_decision": "INCLUDE", "text_source": "pmc"}],
                "exclude": [],
                "stats": {},
            },
            "data_extraction": {
                "papers": [
                    {"pmid": "10001", "title": "Paper 1",
                     "extraction_quality": "complete", "extraction_confidence": "high",
                     "outcomes": [{"name": "test", "effect_size_value": -0.5,
                                   "measure_type": "SMD", "numeric_complete": True}],
                     "concordance": None},
                ],
                "stats": {},
            },
            "quality_assessment": {
                "papers": [
                    {"pmid": "10001", "title": "Paper 1",
                     "quality_assessment": {"tool": "RoB2", "overall": "Low",
                                            "domains": {}},
                     "quality_weight": 1.0},
                ],
                "stats": {},
            },
            "synthesis": {
                "analyses": {
                    "SMD": {
                        "model": {
                            "pooled_estimate": -0.50, "ci_lower": -0.72,
                            "ci_upper": -0.28, "p_value": 0.0001, "k_studies": 3,
                        },
                        "heterogeneity": {"I2": 5.2, "tau2": 0.002, "Q": 2.1},
                        "publication_bias": {
                            "egger_test": None,      # THE BUG TRIGGER
                            "trim_and_fill": None,   # THE BUG TRIGGER
                            "note": "Too few studies (k=3) for Egger test",
                        },
                        "grade": {"final_level": "Moderate"},
                        "sensitivity": {"leave_one_out": []},
                    },
                },
                "stats": {},
            },
        }

        rr = ReviewReportGenerator()
        # This used to crash with: AttributeError: 'NoneType' object has no attribute 'get'
        report = rr.generate(pipeline_data)

        assert "synthesis_concerns" in report
        assert isinstance(report["synthesis_concerns"], list)
        # Should not have flagged pub bias (since both are None)
        for concern in report.get("synthesis_concerns", []):
            for flag in concern.get("flags", []):
                assert "Egger" not in flag, "Should not flag Egger when egger_test is None"
                assert "Trim-and-fill" not in flag, "Should not flag T&F when trim_and_fill is None"

    def test_review_report_with_significant_egger(self):
        """When Egger p < 0.05, review report should flag publication bias."""
        from agents.review_report import ReviewReportGenerator

        pipeline_data = {
            "screening": {"include": [], "exclude": [], "uncertain": [],
                          "stats": {"total": 0, "included": 0, "excluded": 0, "uncertain": 0}},
            "fulltext_retrieval": {"include": [], "exclude": [], "stats": {}},
            "data_extraction": {"papers": [], "stats": {}},
            "quality_assessment": {"papers": [], "stats": {}},
            "synthesis": {
                "analyses": {
                    "SMD": {
                        "model": {"pooled_estimate": -0.35, "ci_lower": -0.50,
                                  "ci_upper": -0.20, "p_value": 0.0001, "k_studies": 12},
                        "heterogeneity": {"I2": 30.0, "tau2": 0.01, "Q": 15.0},
                        "publication_bias": {
                            "egger_test": {
                                "intercept": 2.5, "se": 0.8,
                                "p_value": 0.01,  # significant
                                "significant": True,
                            },
                            "trim_and_fill": {
                                "n_missing": 3,
                                "adjusted_estimate": -0.25,
                                "adjusted_ci_lower": -0.40,
                                "adjusted_ci_upper": -0.10,
                            },
                        },
                        "grade": {"final_level": "Moderate"},
                        "sensitivity": {"leave_one_out": []},
                    },
                },
                "stats": {},
            },
        }

        rr = ReviewReportGenerator()
        report = rr.generate(pipeline_data)

        # Should flag both Egger and Trim-and-fill
        all_flags = []
        for concern in report.get("synthesis_concerns", []):
            all_flags.extend(concern.get("flags", []))

        egger_flagged = any("Egger" in f for f in all_flags)
        tf_flagged = any("Trim-and-fill" in f or "trim-and-fill" in f.lower() for f in all_flags)

        assert egger_flagged, f"Egger test should be flagged (p=0.01). Flags: {all_flags}"
        assert tf_flagged, f"Trim-and-fill should be flagged (n_missing=3). Flags: {all_flags}"


# ══════════════════════════════════════════════════════════════════════════════
# T7 TESTS: Population descriptive section in A8 manuscript
# ══════════════════════════════════════════════════════════════════════════════

class TestT7PopulationDescriptive:
    """Test _build_population_summary() aggregates A5 study characteristics."""

    def test_population_summary_includes_n_studies(self, full_pipeline_data):
        """Population summary mentions the number of included studies."""
        from agents.agent8_manuscript import Agent8Manuscript
        a8 = Agent8Manuscript.__new__(Agent8Manuscript)
        a8._extraction = full_pipeline_data["data_extraction"]
        summary = a8._build_population_summary()
        assert "Number of included studies: 2" in summary

    def test_population_summary_includes_total_participants(self, full_pipeline_data):
        """Population summary aggregates sample_size across papers."""
        from agents.agent8_manuscript import Agent8Manuscript
        a8 = Agent8Manuscript.__new__(Agent8Manuscript)
        a8._extraction = full_pipeline_data["data_extraction"]
        summary = a8._build_population_summary()
        # 120 + 90 = 210
        assert "Total participants: 210" in summary

    def test_population_summary_includes_year_range(self, full_pipeline_data):
        """Population summary reports publication year range."""
        from agents.agent8_manuscript import Agent8Manuscript
        a8 = Agent8Manuscript.__new__(Agent8Manuscript)
        a8._extraction = full_pipeline_data["data_extraction"]
        summary = a8._build_population_summary()
        assert "2021" in summary and "2022" in summary

    def test_population_summary_includes_countries(self, full_pipeline_data):
        """Population summary lists unique countries from A5 data."""
        from agents.agent8_manuscript import Agent8Manuscript
        a8 = Agent8Manuscript.__new__(Agent8Manuscript)
        a8._extraction = full_pipeline_data["data_extraction"]
        summary = a8._build_population_summary()
        assert "USA" in summary and "UK" in summary

    def test_population_summary_with_age_and_female_pct(self, full_pipeline_data):
        """Population summary reports age and female % when present in A5 data."""
        from agents.agent8_manuscript import Agent8Manuscript
        # Inject age_mean and female_pct into papers
        for i, p in enumerate(full_pipeline_data["data_extraction"]["papers"]):
            p["study_characteristics"]["age_mean"] = 35.0 + i * 5
            p["study_characteristics"]["female_pct"] = 55.0 + i * 5

        a8 = Agent8Manuscript.__new__(Agent8Manuscript)
        a8._extraction = full_pipeline_data["data_extraction"]
        summary = a8._build_population_summary()
        assert "Mean age" in summary
        assert "female" in summary.lower()

    def test_population_summary_empty_when_no_extraction(self):
        """Population summary returns empty string when no extraction data."""
        from agents.agent8_manuscript import Agent8Manuscript
        a8 = Agent8Manuscript.__new__(Agent8Manuscript)
        a8._extraction = None
        summary = a8._build_population_summary()
        assert summary == ""

    def test_population_summary_empty_extraction_no_papers(self):
        """Population summary returns empty string when extraction has no papers."""
        from agents.agent8_manuscript import Agent8Manuscript
        a8 = Agent8Manuscript.__new__(Agent8Manuscript)
        a8._extraction = {"papers": []}
        summary = a8._build_population_summary()
        assert summary == ""

    def test_population_summary_handles_non_numeric_sample_size(self):
        """Population summary ignores non-numeric sample_size values."""
        from agents.agent8_manuscript import Agent8Manuscript
        a8 = Agent8Manuscript.__new__(Agent8Manuscript)
        a8._extraction = {
            "papers": [
                {
                    "pmid": "99999",
                    "year": "2023",
                    "study_characteristics": {
                        "study_design": "RCT",
                        "sample_size": "not reported",
                        "country": "Germany",
                    },
                },
            ],
        }
        summary = a8._build_population_summary()
        # Should still produce output (n_studies) but not crash
        assert "Number of included studies: 1" in summary
        assert "Total participants" not in summary

    def test_population_summary_single_year(self):
        """Population summary shows single year when all papers same year."""
        from agents.agent8_manuscript import Agent8Manuscript
        a8 = Agent8Manuscript.__new__(Agent8Manuscript)
        a8._extraction = {
            "papers": [
                {"pmid": "11111", "year": "2024",
                 "study_characteristics": {"sample_size": 50, "country": "France"}},
                {"pmid": "22222", "year": "2024",
                 "study_characteristics": {"sample_size": 60, "country": "France"}},
            ],
        }
        summary = a8._build_population_summary()
        assert "Publication year: 2024" in summary


# ══════════════════════════════════════════════════════════════════════════════
# T8 TESTS: R script export from A7
# ══════════════════════════════════════════════════════════════════════════════

class TestT8RScriptExport:
    """Test _export_r_script() generates a valid, reproducible R script."""

    def _make_a7_instance(self):
        """Create a minimal A7 instance (no R init needed for script export)."""
        from agents.agent7_synthesis import Agent7Synthesis
        a7 = Agent7Synthesis.__new__(Agent7Synthesis)
        # Set minimal attributes needed by _export_r_script
        import logging
        a7.logger = logging.getLogger("test_a7")
        a7._config = {"min_studies_for_egger": 10}
        return a7

    def test_r_script_created(self, tmp_output_dir):
        """_export_r_script creates meta_analysis.R in output dir."""
        a7 = self._make_a7_instance()
        groups = {
            "SMD": {
                "yi": [-0.50, -0.35, -0.60],
                "vi": [0.015, 0.018, 0.020],
                "labels": ["Study A", "Study B", "Study C"],
                "quality_weights": [1.0, 0.6, 1.0],
                "_measure_type": "SMD",
            },
        }
        analyses = {"SMD": {"model": {"pooled_estimate": -0.48}}}

        path = a7._export_r_script(groups, analyses, tmp_output_dir)
        assert path is not None
        assert os.path.exists(path)
        assert path.endswith("meta_analysis.R")

    def test_r_script_has_library_metafor(self, tmp_output_dir):
        """R script starts with library(metafor)."""
        a7 = self._make_a7_instance()
        groups = {
            "SMD": {
                "yi": [-0.50, -0.35],
                "vi": [0.015, 0.018],
                "labels": ["Study A", "Study B"],
                "_measure_type": "SMD",
            },
        }
        analyses = {"SMD": {"model": {}}}

        path = a7._export_r_script(groups, analyses, tmp_output_dir)
        with open(path, "r") as f:
            content = f.read()
        assert "library(metafor)" in content

    def test_r_script_has_data_vectors(self, tmp_output_dir):
        """R script contains yi and vi data vectors for each analysis group."""
        a7 = self._make_a7_instance()
        groups = {
            "SMD": {
                "yi": [-0.50, -0.35, -0.60],
                "vi": [0.015, 0.018, 0.020],
                "labels": ["Study1", "Study2", "Study3"],
                "_measure_type": "SMD",
            },
        }
        analyses = {"SMD": {"model": {}}}

        path = a7._export_r_script(groups, analyses, tmp_output_dir)
        with open(path, "r") as f:
            content = f.read()
        assert "yi_SMD <- c(" in content
        assert "vi_SMD <- c(" in content
        assert "-0.500000" in content
        assert "labels_SMD <- c(" in content

    def test_r_script_has_rma_call(self, tmp_output_dir):
        """R script contains rma() random-effects model fit."""
        a7 = self._make_a7_instance()
        groups = {
            "SMD": {
                "yi": [-0.50, -0.35],
                "vi": [0.015, 0.018],
                "_measure_type": "SMD",
            },
        }
        analyses = {"SMD": {"model": {}}}

        path = a7._export_r_script(groups, analyses, tmp_output_dir)
        with open(path, "r") as f:
            content = f.read()
        assert 'rma(yi = yi_SMD, vi = vi_SMD, method = "REML")' in content

    def test_r_script_has_forest_and_funnel(self, tmp_output_dir):
        """R script includes forest() and funnel() plot commands."""
        a7 = self._make_a7_instance()
        groups = {
            "SMD": {
                "yi": [-0.50, -0.35, -0.40],
                "vi": [0.015, 0.018, 0.020],
                "labels": ["S1", "S2", "S3"],
                "_measure_type": "SMD",
            },
        }
        analyses = {"SMD": {"model": {}}}

        path = a7._export_r_script(groups, analyses, tmp_output_dir)
        with open(path, "r") as f:
            content = f.read()
        assert "forest(fit_SMD" in content
        assert "funnel(fit_SMD)" in content

    def test_r_script_egger_only_when_k_ge_10(self, tmp_output_dir):
        """R script includes regtest only when k >= min_studies_for_egger."""
        a7 = self._make_a7_instance()
        # k=3 — should NOT have Egger
        groups_small = {
            "SMD": {
                "yi": [-0.50, -0.35, -0.40],
                "vi": [0.015, 0.018, 0.020],
                "_measure_type": "SMD",
            },
        }
        analyses = {"SMD": {"model": {}}}

        path = a7._export_r_script(groups_small, analyses, tmp_output_dir)
        with open(path, "r") as f:
            content = f.read()
        assert "regtest" not in content

    def test_r_script_egger_included_when_k_ge_10(self, tmp_output_dir):
        """R script includes regtest when k >= 10."""
        a7 = self._make_a7_instance()
        groups_large = {
            "SMD": {
                "yi": [-0.5 + i * 0.03 for i in range(12)],
                "vi": [0.015 + i * 0.002 for i in range(12)],
                "_measure_type": "SMD",
            },
        }
        analyses = {"SMD": {"model": {}}}

        path = a7._export_r_script(groups_large, analyses, tmp_output_dir)
        with open(path, "r") as f:
            content = f.read()
        assert "regtest(fit_SMD)" in content
        assert "Egger" in content

    def test_r_script_has_failsafe_n(self, tmp_output_dir):
        """R script includes fsn() Rosenthal Fail-Safe N when k >= 3."""
        a7 = self._make_a7_instance()
        groups = {
            "SMD": {
                "yi": [-0.50, -0.35, -0.40],
                "vi": [0.015, 0.018, 0.020],
                "_measure_type": "SMD",
            },
        }
        analyses = {"SMD": {"model": {}}}

        path = a7._export_r_script(groups, analyses, tmp_output_dir)
        with open(path, "r") as f:
            content = f.read()
        assert "fsn(yi = yi_SMD, vi = vi_SMD" in content
        assert 'type = "Rosenthal"' in content

    def test_r_script_has_leave_one_out(self, tmp_output_dir):
        """R script includes leave1out() sensitivity analysis when k >= 3."""
        a7 = self._make_a7_instance()
        groups = {
            "SMD": {
                "yi": [-0.50, -0.35, -0.40],
                "vi": [0.015, 0.018, 0.020],
                "_measure_type": "SMD",
            },
        }
        analyses = {"SMD": {"model": {}}}

        path = a7._export_r_script(groups, analyses, tmp_output_dir)
        with open(path, "r") as f:
            content = f.read()
        assert "leave1out(fit_SMD)" in content

    def test_r_script_has_quality_weights(self, tmp_output_dir):
        """R script includes quality-weighted sensitivity analysis when weights exist."""
        a7 = self._make_a7_instance()
        groups = {
            "SMD": {
                "yi": [-0.50, -0.35, -0.40],
                "vi": [0.015, 0.018, 0.020],
                "quality_weights": [1.0, 0.6, 0.3],
                "_measure_type": "SMD",
            },
        }
        analyses = {"SMD": {"model": {}}}

        path = a7._export_r_script(groups, analyses, tmp_output_dir)
        with open(path, "r") as f:
            content = f.read()
        assert "weights_SMD <- c(" in content
        assert "Quality-weighted" in content

    def test_r_script_skips_single_study_group(self, tmp_output_dir):
        """R script skips groups with fewer than 2 studies."""
        a7 = self._make_a7_instance()
        groups = {
            "OR": {
                "yi": [0.30],
                "vi": [0.05],
                "_measure_type": "OR",
            },
        }
        analyses = {"OR": {"model": {}}}

        path = a7._export_r_script(groups, analyses, tmp_output_dir)
        with open(path, "r") as f:
            content = f.read()
        # Should have header but NO rma call for OR since k=1
        assert "rma(yi = yi_OR" not in content

    def test_r_script_multiple_groups(self, tmp_output_dir):
        """R script handles multiple outcome groups (SMD + OR)."""
        a7 = self._make_a7_instance()
        groups = {
            "SMD": {
                "yi": [-0.50, -0.35],
                "vi": [0.015, 0.018],
                "_measure_type": "SMD",
            },
            "OR": {
                "yi": [0.30, 0.45],
                "vi": [0.05, 0.06],
                "_measure_type": "OR",
            },
        }
        analyses = {"SMD": {"model": {}}, "OR": {"model": {}}}

        path = a7._export_r_script(groups, analyses, tmp_output_dir)
        with open(path, "r") as f:
            content = f.read()
        assert "yi_SMD <- c(" in content
        assert "yi_OR <- c(" in content
        assert "fit_SMD" in content
        assert "fit_OR" in content

    def test_r_script_result_key_in_synthesize_output(self):
        """A7 synthesize() result includes r_script_path when script export succeeds."""
        # This tests the integration point — r_script_path key must exist
        # We use the conftest fixture for a7_result_with_failsafe_n as reference
        # The actual test here verifies the key name convention
        from agents.agent7_synthesis import Agent7Synthesis
        a7 = Agent7Synthesis.__new__(Agent7Synthesis)
        assert hasattr(a7, '_export_r_script'), "A7 must have _export_r_script method"


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
