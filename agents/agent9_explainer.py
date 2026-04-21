"""Agent 9: Explainability & Traceability Report (No LLM)

Post-processing agent that generates:
  1. papers_raw.xlsx — flat Excel with one row per paper, all pipeline decisions
  2. explainability.md — Spanish-language traceability report

No GPU, no LLM — pure Python aggregation of pipeline data.
"""

import os
import time
from datetime import datetime
from typing import Dict, List, Optional

from agents.logger import get_logger

logger = get_logger("Agent9Explainer")


class Agent9Explainer:
    """Generates explainability artifacts from pipeline data.

    Input: full pipeline_data dict (output of A1-A8), run_id, output_dir.
    Output: dict with paths to generated files and summary stats.
    GPU: no -- purely analytical, no LLM calls.
    """

    def run(self, pipeline_data: dict, run_id: str, output_dir: str) -> dict:
        """Generate explainability report and papers_raw Excel.

        Args:
            pipeline_data: Full pipeline output dict (keys: protocol, screening,
                fulltext_retrieval, data_extraction, quality_assessment, synthesis,
                manuscript, prisma, review_report, etc.)
            run_id: Pipeline run identifier (e.g. "run_20260309_190842").
            output_dir: Directory where output files will be written.

        Returns:
            dict with 'excel_path', 'report_path', and 'stats'.
        """
        start = time.time()
        os.makedirs(output_dir, exist_ok=True)

        excel_path = os.path.join(output_dir, "papers_raw.xlsx")
        report_path = os.path.join(output_dir, "explainability.md")

        self._generate_excel(pipeline_data, excel_path)
        self._generate_report(pipeline_data, run_id, report_path)

        elapsed = round(time.time() - start, 2)

        n_all = self._count_all_papers(pipeline_data)
        n_included = len(self._get_included_papers(pipeline_data))
        n_excluded = len(self._get_all_excluded(pipeline_data))

        result = {
            "excel_path": excel_path,
            "report_path": report_path,
            "stats": {
                "total_papers": n_all,
                "included": n_included,
                "excluded": n_excluded,
                "elapsed_s": elapsed,
            },
        }
        logger.info(
            "A9 Explainer complete: %d papers, Excel=%s, Report=%s (%.1fs)",
            n_all, excel_path, report_path, elapsed,
        )
        return result

    # ================================================================== #
    #  EXCEL GENERATION                                                    #
    # ================================================================== #

    def _generate_excel(self, data: dict, output_path: str):
        """Generate papers_raw.xlsx with one row per paper across all stages."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = Workbook()
        bold = Font(bold=True)
        header_fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        wrap = Alignment(wrap_text=True, vertical="top")

        headers = [
            "PMID", "Title", "Authors", "Year", "Journal", "DOI",
            "Abstract",
            "N (sample)", "Population", "Intervention", "Control",
            "Outcome", "Effect Size Type", "Effect Size Value",
            "CI Lower", "CI Upper",
            "A3 Decision", "A3 Confidence",
            "A4 Decision", "A4 Text Source",
            "A5 Extraction Quality", "A5 Confidence",
            "A6 Tool", "A6 Overall", "A6 Weight",
            "Reason Excluded", "Exclusion Stage",
        ]

        ws = wb.active
        ws.title = "All Papers"

        # Write headers
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = wrap

        # Build lookup maps
        extraction_map = self._build_extraction_map(data)
        quality_map = self._build_quality_map(data)

        # Collect all papers with their pipeline journey
        rows = self._collect_all_paper_rows(data, extraction_map, quality_map)

        for i, row in enumerate(rows, 2):
            for col, val in enumerate(row, 1):
                cell = ws.cell(row=i, column=col, value=val)
                if col == 7:  # Abstract column
                    cell.alignment = wrap

        # Auto-width (capped at 60 chars)
        for col_idx in range(1, len(headers) + 1):
            max_len = len(str(ws.cell(row=1, column=col_idx).value or ""))
            for row_idx in range(2, min(len(rows) + 2, 20)):
                val = str(ws.cell(row=row_idx, column=col_idx).value or "")
                max_len = max(max_len, min(len(val), 60))
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max_len + 2

        wb.save(output_path)
        logger.info("papers_raw.xlsx saved: %s (%d rows)", output_path, len(rows))

    def _collect_all_paper_rows(
        self, data: dict, extraction_map: dict, quality_map: dict
    ) -> List[list]:
        """Build flat rows for every paper that entered the pipeline."""
        rows = []

        # A3 screening results (all papers that were screened)
        screening = data.get("screening") or {}
        a3_map = {}  # pmid -> (decision, confidence)
        for bucket in ("include", "exclude", "uncertain"):
            for p in screening.get(bucket, []):
                pmid = str(p.get("pmid", ""))
                decision = p.get("decision", bucket.upper())
                confidence = p.get("confidence", "")
                a3_map[pmid] = (decision, confidence)

        # A4 full-text results
        ft = data.get("fulltext_retrieval") or {}
        a4_map = {}  # pmid -> (decision, text_source)
        for p in ft.get("include", []):
            pmid = str(p.get("pmid", ""))
            a4_map[pmid] = (
                p.get("fulltext_decision", "INCLUDE"),
                p.get("text_source", ""),
            )
        for p in ft.get("exclude", []):
            pmid = str(p.get("pmid", ""))
            a4_map[pmid] = (
                p.get("fulltext_decision", "EXCLUDE"),
                p.get("text_source", ""),
            )
        # Also check 'papers' key (some pipeline outputs use this)
        for p in ft.get("papers", []):
            pmid = str(p.get("pmid", ""))
            if pmid not in a4_map:
                a4_map[pmid] = (
                    p.get("fulltext_decision", ""),
                    p.get("text_source", ""),
                )

        # All PMIDs from A2 search
        search = data.get("search_stats") or {}
        a2_papers = (data.get("screening") or {})  # A3 receives all A2 papers

        # Collect all unique PMIDs in pipeline order
        seen = set()
        ordered_pmids = []

        # From screening (A3 = all papers from A2)
        for bucket in ("include", "exclude", "uncertain"):
            for p in screening.get(bucket, []):
                pmid = str(p.get("pmid", ""))
                if pmid and pmid not in seen:
                    seen.add(pmid)
                    ordered_pmids.append((pmid, p))

        # From A4 (in case some weren't in screening)
        for p in ft.get("include", []) + ft.get("exclude", []) + ft.get("papers", []):
            pmid = str(p.get("pmid", ""))
            if pmid and pmid not in seen:
                seen.add(pmid)
                ordered_pmids.append((pmid, p))

        for pmid, paper in ordered_pmids:
            ext = extraction_map.get(pmid, {})
            qual = quality_map.get(pmid, {})

            # Basic fields
            authors = paper.get("authors", [])
            if isinstance(authors, list):
                authors = "; ".join(authors)

            # Extraction fields
            sc = ext.get("study_characteristics") or {}
            ie = ext.get("intervention_exposure") or ext.get("intervention") or {}
            outcomes_raw = ext.get("outcomes", [])
            oc = outcomes_raw[0] if isinstance(outcomes_raw, list) and outcomes_raw else {}

            # A3 decision
            a3_decision, a3_conf = a3_map.get(pmid, ("", ""))

            # A4 decision
            a4_decision, a4_source = a4_map.get(pmid, ("", ""))

            # Exclusion reason
            reason_excluded = ""
            exclusion_stage = ""
            if a3_decision == "EXCLUDE":
                reason_excluded = paper.get("exclusion_reason", paper.get("reasoning", ""))
                exclusion_stage = "title_abstract"
            elif a4_decision == "EXCLUDE":
                reason_excluded = paper.get("exclusion_reason", paper.get("fulltext_reasoning", ""))
                exclusion_stage = "full_text"

            # Quality assessment
            qa = qual.get("quality_assessment") or {}

            row = [
                pmid,
                paper.get("title", ""),
                authors,
                paper.get("year", ""),
                paper.get("journal", ""),
                paper.get("doi", ""),
                (paper.get("abstract", "") or "")[:500],
                sc.get("sample_size", ""),
                sc.get("population", ""),
                ie.get("name", ie.get("intervention_name", "")),
                ie.get("comparator", ie.get("control", "")),
                oc.get("name", oc.get("outcome", "")),
                oc.get("effect_size_type", oc.get("measure_type", "")),
                oc.get("effect_size_value", oc.get("value", "")),
                oc.get("ci_lower", ""),
                oc.get("ci_upper", ""),
                a3_decision,
                a3_conf,
                a4_decision,
                a4_source,
                ext.get("extraction_quality", ""),
                ext.get("extraction_confidence", ""),
                qa.get("tool", qual.get("tool", "")),
                qa.get("overall", qual.get("overall_judgment", "")),
                qual.get("quality_weight", ""),
                reason_excluded,
                exclusion_stage,
            ]
            rows.append(row)

        return rows

    # ================================================================== #
    #  EXPLAINABILITY REPORT (Spanish)                                     #
    # ================================================================== #

    def _generate_report(self, data: dict, run_id: str, output_path: str):
        """Generate explainability.md in Spanish."""
        lines = []
        now = datetime.now().strftime("%Y-%m-%d %H:%M")

        lines.append("# Informe de Explicabilidad y Trazabilidad")
        lines.append("")
        lines.append(f"**Pipeline:** {run_id}")
        lines.append(f"**Generado:** {now}")
        lines.append("")

        # Section 1: Overview
        self._section_overview(data, lines)

        # Section 2: Traceability per stage
        self._section_traceability(data, lines)

        # Section 3: Exclusion reasons
        self._section_exclusions(data, lines)

        # Section 4: FAISS/RAG traceability
        self._section_rag_traceability(data, lines)

        # Section 5: Borderline decisions
        self._section_borderline(data, lines)

        # Section 6: Data provenance
        self._section_provenance(data, lines)

        with open(output_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))

        logger.info("explainability.md saved: %s", output_path)

    def _section_overview(self, data: dict, lines: list):
        """Resumen general del pipeline."""
        lines.append("## 1. Resumen General")
        lines.append("")

        protocol = data.get("protocol") or {}
        prisma = data.get("prisma") or {}

        lines.append(f"- **Pregunta de investigacion:** {protocol.get('research_question', 'No definida')}")
        lines.append(f"- **Framework:** {protocol.get('framework', 'No definido')}")
        lines.append(f"- **Poblacion:** {protocol.get('population', 'No definida')}")
        lines.append(f"- **Intervencion:** {protocol.get('intervention', protocol.get('exposure', 'No definida'))}")
        lines.append(f"- **Comparador:** {protocol.get('comparator', 'No definido')}")
        lines.append(f"- **Outcome:** {protocol.get('outcome', 'No definido')}")
        lines.append("")

        lines.append("### Flujo PRISMA")
        lines.append("")
        lines.append(f"| Etapa | N |")
        lines.append(f"|-------|---|")
        lines.append(f"| Identificados | {prisma.get('identified', '?')} |")
        lines.append(f"| Tras deduplicacion | {prisma.get('after_dedup', '?')} |")
        lines.append(f"| Cribados (titulo/abstract) | {prisma.get('screened', '?')} |")
        lines.append(f"| Excluidos titulo/abstract | {prisma.get('excluded_title_abstract', '?')} |")
        lines.append(f"| Evaluados texto completo | {prisma.get('assessed_fulltext', '?')} |")
        lines.append(f"| Excluidos texto completo | {prisma.get('excluded_fulltext', '?')} |")
        lines.append(f"| Incluidos final | {prisma.get('included_final', '?')} |")
        lines.append(f"| Datos extraidos | {prisma.get('extracted', '?')} |")
        lines.append(f"| Calidad evaluada | {prisma.get('quality_assessed', '?')} |")
        lines.append("")

    def _section_traceability(self, data: dict, lines: list):
        """Trazabilidad: origen de cada dato por etapa."""
        lines.append("## 2. Trazabilidad por Etapa")
        lines.append("")

        # A1
        protocol = data.get("protocol") or {}
        lines.append("### A1 — Protocolo")
        lines.append("")
        lines.append(f"- Query PubMed generada: `{protocol.get('pubmed_query', data.get('pubmed_query', 'N/A'))}`")
        inc = protocol.get("inclusion_criteria", [])
        if inc:
            lines.append("- Criterios de inclusion:")
            for c in inc:
                lines.append(f"  - {c}")
        exc = protocol.get("exclusion_criteria", [])
        if exc:
            lines.append("- Criterios de exclusion:")
            for c in exc:
                lines.append(f"  - {c}")
        lines.append("")

        # A2
        search_stats = data.get("search_stats") or {}
        lines.append("### A2 — Busqueda")
        lines.append("")
        lines.append(f"- PubMed: {search_stats.get('pubmed', 0)} resultados")
        lines.append(f"- Europe PMC: {search_stats.get('europepmc', 0)} resultados")
        lines.append(f"- Semantic Scholar: {search_stats.get('semantic_scholar', 0)} resultados")
        lines.append(f"- Google Scholar: {search_stats.get('google_scholar', 0)} resultados")
        lines.append(f"- Total antes de deduplicacion: {search_stats.get('total_before_dedup', 0)}")
        lines.append(f"- Duplicados eliminados: {search_stats.get('duplicates_removed', 0)}")
        lines.append(f"- Total tras deduplicacion: {search_stats.get('total_after_dedup', 0)}")
        lines.append("")

        # A3
        screening = data.get("screening") or {}
        a3_stats = screening.get("stats") or {}
        lines.append("### A3 — Cribado Titulo/Abstract")
        lines.append("")
        lines.append(f"- Total cribados: {a3_stats.get('total', 0)}")
        lines.append(f"- Incluidos: {a3_stats.get('included', 0)}")
        lines.append(f"- Excluidos: {a3_stats.get('excluded', 0)}")
        lines.append(f"- Inciertos: {a3_stats.get('uncertain', 0)}")
        lines.append(f"- Ensemble size: {a3_stats.get('ensemble_size', '?')}")
        lines.append(f"- Umbral de votacion: {a3_stats.get('voting_threshold', '?')}")
        lines.append("")

        # A4
        ft = data.get("fulltext_retrieval") or {}
        a4_stats = ft.get("stats") or {}
        lines.append("### A4 — Texto Completo")
        lines.append("")
        lines.append(f"- PMC: {a4_stats.get('pmc_fulltext', 0)}")
        lines.append(f"- Unpaywall: {a4_stats.get('unpaywall_pdf', 0)}")
        lines.append(f"- Solo abstract: {a4_stats.get('abstract_only', 0)}")
        lines.append(f"- Incluidos tras texto completo: {a4_stats.get('fulltext_included', 0)}")
        lines.append(f"- Excluidos tras texto completo: {a4_stats.get('fulltext_excluded', 0)}")
        lines.append("")

        # A5
        ext = data.get("data_extraction") or {}
        a5_stats = ext.get("stats") or {}
        lines.append("### A5 — Extraccion de Datos")
        lines.append("")
        lines.append(f"- Extraccion completa: {a5_stats.get('complete_extraction', 0)}")
        lines.append(f"- Extraccion parcial: {a5_stats.get('partial_extraction', 0)}")
        lines.append(f"- Extraccion fallida: {a5_stats.get('failed_extraction', 0)}")
        lines.append(f"- Datos numericos completos: {a5_stats.get('numeric_complete', 0)}")
        lines.append(f"- Recuperados (4o pase): {a5_stats.get('numeric_recovered', 0)}")
        lines.append(f"- Extraccion dual: {'Si' if a5_stats.get('dual_extraction') else 'No'}")
        lines.append("")

        # A6
        qa = data.get("quality_assessment") or {}
        a6_stats = qa.get("stats") or {}
        lines.append("### A6 — Evaluacion de Calidad")
        lines.append("")
        lines.append(f"- RoB 2 (ECAs): {a6_stats.get('rob2_assessed', 0)}")
        lines.append(f"- ROBINS-I (observacionales): {a6_stats.get('robins_i_assessed', 0)}")
        lines.append(f"- JBI (cualitativos): {a6_stats.get('jbi_assessed', 0)}")
        lines.append(f"- Fallidos: {a6_stats.get('failed', 0)}")
        lines.append("")

        # A7
        synthesis = data.get("synthesis") or {}
        a7_stats = synthesis.get("stats") or {}
        analyses = synthesis.get("analyses") or {}
        lines.append("### A7 — Sintesis Estadistica")
        lines.append("")
        lines.append(f"- Papers con datos utilizables: {a7_stats.get('papers_with_usable_data', 0)}")
        lines.append(f"- Analisis realizados: {a7_stats.get('analyses_performed', 0)}")
        lines.append(f"- Medidas encontradas: {', '.join(a7_stats.get('measure_types_found', []))}")
        for measure, analysis in analyses.items():
            model = analysis.get("model") or {}
            het = analysis.get("heterogeneity") or {}
            grade = analysis.get("grade") or {}
            lines.append(f"  - **{measure}**: estimacion={model.get('pooled_estimate', '?')}, "
                         f"IC 95% [{model.get('ci_lower', '?')}, {model.get('ci_upper', '?')}], "
                         f"I2={het.get('I2', model.get('I2', '?'))}%, "
                         f"GRADE={grade.get('final_level', '?')}")
        lines.append("")

        # A8
        manuscript = data.get("manuscript") or {}
        a8_stats = manuscript.get("stats") or {}
        lines.append("### A8 — Manuscrito")
        lines.append("")
        lines.append(f"- Chunks RAG totales: {a8_stats.get('total_chunks', '?')}")
        lines.append(f"- Papers citados: {a8_stats.get('total_papers_cited', '?')}")
        cps = a8_stats.get("chunks_per_section") or {}
        if cps:
            lines.append("- Chunks por seccion:")
            for sec, n in cps.items():
                lines.append(f"  - {sec}: {n}")
        lines.append(f"- PDF generado: {'Si' if manuscript.get('pdf_path') else 'No'}")
        lines.append("")

    def _section_exclusions(self, data: dict, lines: list):
        """Razon de exclusion por paper."""
        lines.append("## 3. Razones de Exclusion por Paper")
        lines.append("")

        excluded = self._get_all_excluded(data)
        if not excluded:
            lines.append("No se excluyeron papers.")
            lines.append("")
            return

        lines.append(f"| PMID | Titulo | Etapa | Razon |")
        lines.append(f"|------|--------|-------|-------|")
        for p in excluded:
            title = (p.get("title", "") or "")[:80]
            # Escape pipes in title
            title = title.replace("|", "/")
            reason = (p.get("reason", "") or "").replace("|", "/")
            lines.append(f"| {p.get('pmid', '')} | {title} | {p.get('stage', '')} | {reason} |")
        lines.append("")

    def _section_rag_traceability(self, data: dict, lines: list):
        """Que chunks de FAISS alimentaron cada seccion del manuscrito."""
        lines.append("## 4. Trazabilidad RAG (FAISS -> Manuscrito)")
        lines.append("")

        manuscript = data.get("manuscript") or {}
        a8_stats = manuscript.get("stats") or {}
        cps = a8_stats.get("chunks_per_section") or {}

        if not cps:
            lines.append("No hay datos de trazabilidad RAG disponibles.")
            lines.append("")
            return

        total_chunks = a8_stats.get("total_chunks", 0)
        lines.append(f"El indice FAISS contiene **{total_chunks} chunks** de texto de los papers incluidos.")
        lines.append("")
        lines.append("Cada seccion del manuscrito se genero consultando los chunks mas relevantes:")
        lines.append("")
        lines.append("| Seccion | Chunks usados |")
        lines.append("|---------|---------------|")
        for sec, n in cps.items():
            lines.append(f"| {sec} | {n} |")
        lines.append("")
        lines.append(
            "Los chunks se seleccionan por similitud coseno (MiniLM-L6-v2) con una query "
            "especifica por seccion. Cada chunk incluye metadatos de PMID y seccion de origen, "
            "permitiendo trazar que papers informaron cada parte del manuscrito."
        )
        lines.append("")

        # Citation verification
        cv = manuscript.get("citation_verification") or {}
        if cv:
            lines.append("### Verificacion de citas")
            lines.append("")
            lines.append(f"- Citas totales: {cv.get('total_citations', '?')}")
            lines.append(f"- Verificadas: {cv.get('total_verified', '?')}")
            lines.append(f"- Citas fantasma eliminadas: {cv.get('total_removed', 0)}")
            lines.append("")

    def _section_borderline(self, data: dict, lines: list):
        """Decisiones borderline del review report."""
        lines.append("## 5. Decisiones Borderline")
        lines.append("")

        review_report = data.get("review_report") or {}
        if not review_report:
            # Try generating from ReviewReportGenerator
            try:
                from agents.review_report import ReviewReportGenerator
                rr = ReviewReportGenerator()
                review_report = rr.generate(data)
            except Exception as e:
                logger.warning("Could not generate review report: %s", e)
                lines.append("No se pudo generar el analisis de decisiones borderline.")
                lines.append("")
                return

        summary = review_report.get("summary") or {}
        total_flags = summary.get("total_flags", 0)

        if total_flags == 0:
            lines.append("No se identificaron decisiones borderline en el pipeline.")
            lines.append("")
            return

        lines.append(f"Se identificaron **{total_flags} alertas** en el pipeline:")
        lines.append(f"- Prioridad alta: {summary.get('high_priority', 0)}")
        lines.append(f"- Prioridad media: {summary.get('medium_priority', 0)}")
        lines.append(f"- Prioridad baja: {summary.get('low_priority', 0)}")
        lines.append(f"- Confianza del pipeline: {summary.get('pipeline_confidence', '?')}")
        lines.append("")

        # Screening borderline
        scr = review_report.get("screening_borderline", [])
        if scr:
            lines.append("### Cribado borderline (A3)")
            lines.append("")
            for item in scr:
                lines.append(f"- **PMID {item.get('pmid', '?')}** [{item.get('priority', '?').upper()}]: "
                             f"{item.get('title', '?')[:70]}")
                for flag in item.get("flags", []):
                    lines.append(f"  - {flag}")
            lines.append("")

        # Extraction concerns
        ext_c = review_report.get("extraction_concerns", [])
        if ext_c:
            lines.append("### Problemas de extraccion (A5)")
            lines.append("")
            for item in ext_c:
                lines.append(f"- **PMID {item.get('pmid', '?')}** [{item.get('priority', '?').upper()}]: "
                             f"{item.get('title', '?')[:70]}")
                for flag in item.get("flags", []):
                    lines.append(f"  - {flag}")
            lines.append("")

        # Quality concerns
        qa_c = review_report.get("quality_concerns", [])
        if qa_c:
            lines.append("### Problemas de calidad (A6)")
            lines.append("")
            for item in qa_c:
                lines.append(f"- **PMID {item.get('pmid', '?')}** [{item.get('priority', '?').upper()}]: "
                             f"{item.get('title', '?')[:70]}")
                for flag in item.get("flags", []):
                    lines.append(f"  - {flag}")
            lines.append("")

        # Synthesis concerns
        syn_c = review_report.get("synthesis_concerns", [])
        if syn_c:
            lines.append("### Problemas de sintesis (A7)")
            lines.append("")
            for item in syn_c:
                lines.append(f"- **{item.get('measure', '?')}** [{item.get('priority', '?').upper()}]")
                for flag in item.get("flags", []):
                    lines.append(f"  - {flag}")
            lines.append("")

        # Recommendations
        recs = review_report.get("recommendations", [])
        if recs:
            lines.append("### Recomendaciones")
            lines.append("")
            for i, rec in enumerate(recs, 1):
                lines.append(f"{i}. {rec}")
            lines.append("")

    def _section_provenance(self, data: dict, lines: list):
        """Origen de los datos de cada paper incluido."""
        lines.append("## 6. Procedencia de Datos por Paper Incluido")
        lines.append("")

        included = self._get_included_papers(data)
        extraction_map = self._build_extraction_map(data)
        quality_map = self._build_quality_map(data)

        # Databases A2 actually queried, propagated by the orchestrator
        # (api.py::run_pipeline_generator). Falls back to deriving from
        # search_stats if missing, and finally to "(no disponible)".
        a2_sources = data.get("a2_sources") or []
        if not a2_sources:
            search_stats = data.get("search_stats") or {}
            a2_sources = [
                db for db in ("pubmed", "europepmc", "semantic_scholar", "google_scholar")
                if search_stats.get(db, 0) > 0
            ]
        if a2_sources:
            lines.append(
                "**Bases de datos consultadas por A2:** "
                + ", ".join(a2_sources)
            )
        else:
            lines.append("**Bases de datos consultadas por A2:** (no disponible)")
        lines.append("")

        if not included:
            lines.append("No hay papers incluidos.")
            lines.append("")
            return

        for p in included:
            pmid = str(p.get("pmid", ""))
            title = p.get("title", "Sin titulo")
            lines.append(f"### PMID {pmid}")
            lines.append(f"**{title}**")
            lines.append("")

            # Source — per-paper `source` (singular) or merged `sources` list.
            # If both are absent we surface "(no disponible)" instead of "?".
            src = p.get("source")
            if not src:
                src_list = p.get("sources") or []
                if isinstance(src_list, list) and src_list:
                    src = ", ".join(src_list)
            lines.append(f"- Fuente de busqueda: {src or '(no disponible)'}")
            lines.append(f"- Texto completo: {p.get('text_source', '(no disponible)')}")

            # Extraction
            ext = extraction_map.get(pmid, {})
            if ext:
                lines.append(f"- Calidad de extraccion: {ext.get('extraction_quality', '?')}")
                lines.append(f"- Confianza: {ext.get('extraction_confidence', '?')}")
                sections_used = ext.get("sections_used", [])
                if sections_used:
                    lines.append(f"- Secciones usadas: {', '.join(sections_used)}")
                outcomes = ext.get("outcomes", [])
                for oc in outcomes:
                    name = oc.get("name", "?")
                    val = oc.get("value", oc.get("effect_size_value", "?"))
                    mt = oc.get("measure_type", "?")
                    nc = "si" if oc.get("numeric_complete") else "no"
                    lines.append(f"  - Outcome: {name} ({mt}={val}, numerico completo: {nc})")

            # Quality
            qual = quality_map.get(pmid, {})
            if qual:
                qa = qual.get("quality_assessment") or {}
                lines.append(f"- Herramienta calidad: {qa.get('tool', qual.get('tool', '?'))}")
                lines.append(f"- Juicio global: {qa.get('overall', qual.get('overall_judgment', '?'))}")
                lines.append(f"- Peso de calidad: {qual.get('quality_weight', '?')}")

            lines.append("")

    # ================================================================== #
    #  HELPER METHODS                                                      #
    # ================================================================== #

    @staticmethod
    def _get_included_papers(data: dict) -> list:
        """Return included papers from A4."""
        ft = data.get("fulltext_retrieval") or {}
        return ft.get("include", [])

    @staticmethod
    def _get_all_excluded(data: dict) -> list:
        """Return all excluded papers across A3 and A4."""
        excluded = []
        screening = data.get("screening") or {}
        for p in screening.get("exclude", []):
            excluded.append({
                "pmid": p.get("pmid", ""),
                "title": p.get("title", ""),
                "stage": "title_abstract",
                "reason": p.get("exclusion_reason", p.get("reasoning", "")),
            })
        ft = data.get("fulltext_retrieval") or {}
        for p in ft.get("exclude", []):
            excluded.append({
                "pmid": p.get("pmid", ""),
                "title": p.get("title", ""),
                "stage": "full_text",
                "reason": p.get("exclusion_reason", p.get("fulltext_reasoning", "")),
            })
        return excluded

    @staticmethod
    def _count_all_papers(data: dict) -> int:
        """Count total unique papers that entered the pipeline."""
        pmids = set()
        screening = data.get("screening") or {}
        for bucket in ("include", "exclude", "uncertain"):
            for p in screening.get(bucket, []):
                pmids.add(str(p.get("pmid", "")))
        return len(pmids) or 0

    @staticmethod
    def _build_extraction_map(data: dict) -> dict:
        """PMID -> extraction data dict."""
        ext = data.get("data_extraction") or {}
        result = {}
        for p in ext.get("papers", []):
            pmid = str(p.get("pmid", ""))
            if pmid:
                result[pmid] = p
        return result

    @staticmethod
    def _build_quality_map(data: dict) -> dict:
        """PMID -> quality assessment data dict."""
        qa = data.get("quality_assessment") or {}
        result = {}
        for p in qa.get("papers", []):
            pmid = str(p.get("pmid", ""))
            if pmid:
                result[pmid] = p
        return result
