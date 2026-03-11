"""Export functions for SAM pipeline data.

Provides CSV, Excel (openpyxl), RevMan XML, and summary table (Markdown) exports.
All functions accept the full pipeline_data dict and an output_path, returning the path.
"""

import csv
import io
import os
import time
import xml.etree.ElementTree as ET
from xml.dom import minidom
from typing import Optional

from agents.logger import get_logger

logger = get_logger("Exporters")


# ══════════════════════════════════════════════════════════════════════════════
# Helpers
# ══════════════════════════════════════════════════════════════════════════════

def _get_included_papers(data):
    """Return list of included papers from A4 results."""
    a4 = data.get("fulltext_retrieval") or {}
    return a4.get("include", [])


def _get_excluded_papers(data):
    """Return list of excluded papers (A3 + A4)."""
    excluded = []
    screening = data.get("screening") or {}
    for p in screening.get("exclude", []):
        excluded.append({**p, "stage": "title_abstract"})
    a4 = data.get("fulltext_retrieval") or {}
    for p in a4.get("exclude", []):
        excluded.append({**p, "stage": "full_text"})
    return excluded


def _get_extraction_map(data):
    """Return dict mapping PMID -> extraction data."""
    a5 = data.get("data_extraction") or {}
    result = {}
    for p in a5.get("papers", []):
        pmid = str(p.get("pmid", ""))
        if pmid:
            result[pmid] = p
    return result


def _get_quality_map(data):
    """Return dict mapping PMID -> quality assessment data."""
    a6 = data.get("quality_assessment") or {}
    result = {}
    for p in a6.get("papers", []):
        pmid = str(p.get("pmid", ""))
        if pmid:
            result[pmid] = p
    return result


def _safe_get(d, *keys, default=""):
    """Safely traverse nested dicts."""
    current = d
    for k in keys:
        if isinstance(current, dict):
            current = current.get(k, default)
        else:
            return default
    return current if current is not None else default


# ══════════════════════════════════════════════════════════════════════════════
# 1. CSV Export
# ══════════════════════════════════════════════════════════════════════════════

def export_csv(pipeline_data, output_path):
    """Create a comprehensive CSV with one row per included study.

    Args:
        pipeline_data: Full pipeline output dict.
        output_path: Path for the CSV file.

    Returns:
        str: The output path.
    """
    papers = _get_included_papers(pipeline_data)
    extraction_map = _get_extraction_map(pipeline_data)
    quality_map = _get_quality_map(pipeline_data)

    headers = [
        "PMID", "DOI", "Title", "Authors", "Year", "Journal",
        "Study Design", "Intervention", "Comparator", "Population",
        "Outcome", "Effect Size Type", "Effect Size Value",
        "CI Lower", "CI Upper", "P Value",
        "N Intervention", "N Control",
        "RoB Tool", "RoB Overall", "Quality Weight",
        "Screening Decision", "Full-text Source",
    ]

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

    with open(output_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(headers)

        for p in papers:
            pmid = str(p.get("pmid", ""))
            ext = extraction_map.get(pmid, {})
            qual = quality_map.get(pmid, {})

            # Extraction data lives in sub-dicts
            study_chars = ext.get("study_characteristics", {})
            intervention = ext.get("intervention_exposure", ext.get("intervention", {}))
            outcomes_raw = ext.get("outcomes", [])

            # outcomes may be a list (A5 format) or a dict (legacy)
            if isinstance(outcomes_raw, list):
                outcomes = outcomes_raw[0] if outcomes_raw else {}
            elif isinstance(outcomes_raw, dict):
                outcomes = outcomes_raw
            else:
                outcomes = {}

            # Some extractions store these at top level
            if not study_chars and ext.get("design"):
                study_chars = ext

            authors = p.get("authors", [])
            if isinstance(authors, list):
                authors = "; ".join(authors)

            row = [
                pmid,
                p.get("doi", ""),
                p.get("title", ""),
                authors,
                p.get("year", study_chars.get("year", "")),
                p.get("journal", study_chars.get("journal", "")),
                study_chars.get("design", study_chars.get("study_design", "")),
                intervention.get("name", intervention.get("intervention_name", "")),
                intervention.get("comparator", intervention.get("control", "")),
                study_chars.get("population", ""),
                outcomes.get("name", outcomes.get("primary_outcome", outcomes.get("outcome", ""))),
                outcomes.get("effect_size_type", outcomes.get("measure_type", "")),
                outcomes.get("effect_size_value", outcomes.get("effect_size", "")),
                outcomes.get("ci_lower", ""),
                outcomes.get("ci_upper", ""),
                outcomes.get("p_value", ""),
                outcomes.get("n_intervention", outcomes.get("n_treatment", "")),
                outcomes.get("n_control", ""),
                qual.get("tool", ""),
                qual.get("overall_judgment", qual.get("overall", "")),
                qual.get("quality_weight", ""),
                "INCLUDE",
                p.get("text_source", ""),
            ]
            writer.writerow(row)

    logger.info("CSV exported: %s (%d studies)", output_path, len(papers))
    return output_path


# ══════════════════════════════════════════════════════════════════════════════
# 2. Excel Export (multi-sheet)
# ══════════════════════════════════════════════════════════════════════════════

def export_excel(pipeline_data, output_path):
    """Create a multi-sheet Excel workbook.

    Sheets: Included Studies, Excluded Studies, Extraction Data,
            Quality Assessment, Synthesis Results, GRADE Summary.

    Args:
        pipeline_data: Full pipeline output dict.
        output_path: Path for the .xlsx file.

    Returns:
        str: The output path.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    bold = Font(bold=True)

    def write_headers(ws, headers):
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = bold

    # ── Sheet 1: Included Studies ──
    ws1 = wb.active
    ws1.title = "Included Studies"
    h1 = ["PMID", "DOI", "Title", "Authors", "Year", "Journal",
          "Text Source", "Full-text Length", "Screening Reasoning"]
    write_headers(ws1, h1)
    for i, p in enumerate(_get_included_papers(pipeline_data), 2):
        authors = p.get("authors", [])
        if isinstance(authors, list):
            authors = "; ".join(authors)
        ws1.cell(row=i, column=1, value=str(p.get("pmid", "")))
        ws1.cell(row=i, column=2, value=p.get("doi", ""))
        ws1.cell(row=i, column=3, value=p.get("title", ""))
        ws1.cell(row=i, column=4, value=authors)
        ws1.cell(row=i, column=5, value=p.get("year", ""))
        ws1.cell(row=i, column=6, value=p.get("journal", ""))
        ws1.cell(row=i, column=7, value=p.get("text_source", ""))
        ws1.cell(row=i, column=8, value=p.get("full_text_length", 0))
        ws1.cell(row=i, column=9, value=p.get("fulltext_reasoning", ""))

    # ── Sheet 2: Excluded Studies ──
    ws2 = wb.create_sheet("Excluded Studies")
    h2 = ["PMID", "Title", "Stage", "Reason"]
    write_headers(ws2, h2)
    for i, p in enumerate(_get_excluded_papers(pipeline_data), 2):
        ws2.cell(row=i, column=1, value=str(p.get("pmid", "")))
        ws2.cell(row=i, column=2, value=p.get("title", ""))
        ws2.cell(row=i, column=3, value=p.get("stage", ""))
        ws2.cell(row=i, column=4, value=p.get("exclusion_reason", p.get("reason", "")))

    # ── Sheet 3: Extraction Data ──
    ws3 = wb.create_sheet("Extraction Data")
    h3 = ["PMID", "Title", "Design", "Population", "Sample Size",
          "Intervention", "Comparator", "Outcome", "Effect Size Type",
          "Effect Size", "CI Lower", "CI Upper", "P Value",
          "N Intervention", "N Control", "Extraction Quality", "Confidence"]
    write_headers(ws3, h3)
    ext_data = (pipeline_data.get("data_extraction") or {}).get("papers", [])
    for i, p in enumerate(ext_data, 2):
        sc = p.get("study_characteristics", {})
        iv = p.get("intervention_exposure", p.get("intervention", {}))
        oc_raw = p.get("outcomes", [])
        oc = oc_raw[0] if isinstance(oc_raw, list) and oc_raw else (oc_raw if isinstance(oc_raw, dict) else {})
        ws3.cell(row=i, column=1, value=str(p.get("pmid", "")))
        ws3.cell(row=i, column=2, value=p.get("title", ""))
        ws3.cell(row=i, column=3, value=sc.get("design", sc.get("study_design", "")))
        ws3.cell(row=i, column=4, value=sc.get("population", ""))
        ws3.cell(row=i, column=5, value=sc.get("sample_size", sc.get("n", "")))
        ws3.cell(row=i, column=6, value=iv.get("name", iv.get("intervention_name", "")))
        ws3.cell(row=i, column=7, value=iv.get("comparator", iv.get("control", "")))
        ws3.cell(row=i, column=8, value=oc.get("name", oc.get("primary_outcome", oc.get("outcome", ""))))
        ws3.cell(row=i, column=9, value=oc.get("effect_size_type", oc.get("measure_type", "")))
        ws3.cell(row=i, column=10, value=oc.get("effect_size_value", oc.get("effect_size", "")))
        ws3.cell(row=i, column=11, value=oc.get("ci_lower", ""))
        ws3.cell(row=i, column=12, value=oc.get("ci_upper", ""))
        ws3.cell(row=i, column=13, value=oc.get("p_value", ""))
        ws3.cell(row=i, column=14, value=oc.get("n_intervention", oc.get("n_treatment", "")))
        ws3.cell(row=i, column=15, value=oc.get("n_control", ""))
        ws3.cell(row=i, column=16, value=p.get("extraction_quality", ""))
        ws3.cell(row=i, column=17, value=p.get("extraction_confidence", ""))

    # ── Sheet 4: Quality Assessment ──
    ws4 = wb.create_sheet("Quality Assessment")
    h4 = ["PMID", "Title", "Study Design", "Tool", "Overall Judgment",
          "Quality Weight", "Domains"]
    write_headers(ws4, h4)
    qa_data = (pipeline_data.get("quality_assessment") or {}).get("papers", [])
    for i, p in enumerate(qa_data, 2):
        domains = p.get("domains", {})
        domain_str = "; ".join(f"{k}: {v}" for k, v in domains.items()) if isinstance(domains, dict) else str(domains)
        ws4.cell(row=i, column=1, value=str(p.get("pmid", "")))
        ws4.cell(row=i, column=2, value=p.get("title", ""))
        ws4.cell(row=i, column=3, value=p.get("study_design", ""))
        ws4.cell(row=i, column=4, value=p.get("tool", ""))
        ws4.cell(row=i, column=5, value=p.get("overall_judgment", p.get("overall", "")))
        ws4.cell(row=i, column=6, value=p.get("quality_weight", ""))
        ws4.cell(row=i, column=7, value=domain_str)

    # ── Sheet 5: Synthesis Results ──
    ws5 = wb.create_sheet("Synthesis Results")
    h5 = ["Measure Type", "Pooled Estimate", "CI Lower", "CI Upper",
          "P Value", "I-squared (%)", "Tau-squared", "Q Statistic",
          "Q P-value", "K Studies", "Method"]
    write_headers(ws5, h5)
    analyses = (pipeline_data.get("synthesis") or {}).get("analyses", {})
    row = 2
    for measure, analysis in analyses.items():
        model = analysis.get("model", {})
        het = analysis.get("heterogeneity", {})
        ws5.cell(row=row, column=1, value=measure)
        ws5.cell(row=row, column=2, value=model.get("pooled_estimate", ""))
        ws5.cell(row=row, column=3, value=model.get("ci_lower", ""))
        ws5.cell(row=row, column=4, value=model.get("ci_upper", ""))
        ws5.cell(row=row, column=5, value=model.get("p_value", ""))
        ws5.cell(row=row, column=6, value=model.get("I2", het.get("I2", "")))
        ws5.cell(row=row, column=7, value=het.get("tau2", ""))
        ws5.cell(row=row, column=8, value=het.get("Q", ""))
        ws5.cell(row=row, column=9, value=het.get("Q_pvalue", ""))
        ws5.cell(row=row, column=10, value=model.get("k", ""))
        ws5.cell(row=row, column=11, value=model.get("method", "REML"))
        row += 1

    # ── Sheet 6: GRADE Summary ──
    ws6 = wb.create_sheet("GRADE Summary")
    h6 = ["Measure Type", "Starting Level", "Final Level",
          "Risk of Bias", "Inconsistency", "Indirectness",
          "Imprecision", "Publication Bias", "Downgrades", "Rationale"]
    write_headers(ws6, h6)
    row = 2
    for measure, analysis in analyses.items():
        grade = analysis.get("grade", {})
        if not grade:
            continue
        domains = grade.get("domains", {})
        downgrades = grade.get("downgrades", [])
        ws6.cell(row=row, column=1, value=measure)
        ws6.cell(row=row, column=2, value=grade.get("starting_level", ""))
        ws6.cell(row=row, column=3, value=grade.get("final_level", ""))
        ws6.cell(row=row, column=4, value=domains.get("risk_of_bias", ""))
        ws6.cell(row=row, column=5, value=domains.get("inconsistency", ""))
        ws6.cell(row=row, column=6, value=domains.get("indirectness", ""))
        ws6.cell(row=row, column=7, value=domains.get("imprecision", ""))
        ws6.cell(row=row, column=8, value=domains.get("publication_bias", ""))
        ws6.cell(row=row, column=9, value="; ".join(downgrades) if isinstance(downgrades, list) else str(downgrades))
        ws6.cell(row=row, column=10, value=grade.get("rationale", ""))
        row += 1

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    wb.save(output_path)
    logger.info("Excel exported: %s", output_path)
    return output_path


# ══════════════════════════════════════════════════════════════════════════════
# 3. RevMan 5 XML Export
# ══════════════════════════════════════════════════════════════════════════════

def export_revman(pipeline_data, output_path):
    """Generate a simplified RevMan 5 compatible XML file.

    Args:
        pipeline_data: Full pipeline output dict.
        output_path: Path for the .xml file.

    Returns:
        str: The output path.
    """
    protocol = pipeline_data.get("protocol", {})
    papers = _get_included_papers(pipeline_data)
    extraction_map = _get_extraction_map(pipeline_data)
    quality_map = _get_quality_map(pipeline_data)

    # Root element
    root = ET.Element("COCHRANE_REVIEW")
    root.set("REVMAN_SUB_VERSION", "5.4")
    root.set("MODIFIED", time.strftime("%Y-%m-%d"))

    # Cover sheet
    cover = ET.SubElement(root, "COVER_SHEET")
    title_el = ET.SubElement(cover, "TITLE")
    title_el.text = protocol.get("research_question", "Systematic Review")
    ET.SubElement(cover, "REVIEW_TYPE").text = "Intervention"

    # Characteristics of included studies
    studies_el = ET.SubElement(root, "CHARACTERISTICS_OF_STUDIES")
    included_el = ET.SubElement(studies_el, "CHARACTERISTICS_OF_INCLUDED_STUDIES")

    for p in papers:
        pmid = str(p.get("pmid", ""))
        ext = extraction_map.get(pmid, {})
        qual = quality_map.get(pmid, {})

        study_el = ET.SubElement(included_el, "INCLUDED_STUDY")
        study_el.set("STUDY_ID", pmid)

        # Study identification
        authors = p.get("authors", [])
        first_author = authors[0] if authors else "Unknown"
        year = str(p.get("year", ""))
        study_el.set("NAME", f"{first_author} {year}")

        # Characteristics
        sc = ext.get("study_characteristics", {})
        iv = ext.get("intervention_exposure", ext.get("intervention", {}))
        oc_raw = ext.get("outcomes", [])
        oc = oc_raw[0] if isinstance(oc_raw, list) and oc_raw else (oc_raw if isinstance(oc_raw, dict) else {})

        chars_el = ET.SubElement(study_el, "CHAR_METHODS")
        ET.SubElement(chars_el, "CHAR_DESIGN").text = sc.get("design", sc.get("study_design", ""))
        ET.SubElement(chars_el, "CHAR_PARTICIPANTS").text = sc.get("population", "")
        ET.SubElement(chars_el, "CHAR_SAMPLE_SIZE").text = str(sc.get("sample_size", sc.get("n", "")))
        ET.SubElement(chars_el, "CHAR_SETTING").text = sc.get("setting", "")

        intv_el = ET.SubElement(study_el, "CHAR_INTERVENTIONS")
        ET.SubElement(intv_el, "INTERVENTION").text = iv.get("name", iv.get("intervention_name", ""))
        ET.SubElement(intv_el, "COMPARATOR").text = iv.get("comparator", iv.get("control", ""))
        ET.SubElement(intv_el, "DOSE").text = iv.get("dose", iv.get("dosage", ""))
        ET.SubElement(intv_el, "DURATION").text = iv.get("duration", "")

        outcomes_el = ET.SubElement(study_el, "CHAR_OUTCOMES")
        ET.SubElement(outcomes_el, "OUTCOME").text = oc.get("name", oc.get("primary_outcome", oc.get("outcome", "")))
        ET.SubElement(outcomes_el, "EFFECT_SIZE_TYPE").text = oc.get("effect_size_type", oc.get("measure_type", ""))
        ET.SubElement(outcomes_el, "EFFECT_SIZE").text = str(oc.get("effect_size_value", oc.get("effect_size", "")))
        ET.SubElement(outcomes_el, "CI_LOWER").text = str(oc.get("ci_lower", ""))
        ET.SubElement(outcomes_el, "CI_UPPER").text = str(oc.get("ci_upper", ""))
        ET.SubElement(outcomes_el, "P_VALUE").text = str(oc.get("p_value", ""))

        # Risk of bias
        rob_el = ET.SubElement(study_el, "RISK_OF_BIAS")
        ET.SubElement(rob_el, "TOOL").text = qual.get("tool", "")
        ET.SubElement(rob_el, "OVERALL").text = qual.get("overall_judgment", qual.get("overall", ""))
        ET.SubElement(rob_el, "WEIGHT").text = str(qual.get("quality_weight", ""))
        domains = qual.get("domains", {})
        if isinstance(domains, dict):
            for domain_name, judgment in domains.items():
                d_el = ET.SubElement(rob_el, "DOMAIN")
                d_el.set("NAME", domain_name)
                d_el.text = str(judgment)

    # Quality of evidence (GRADE)
    analyses = (pipeline_data.get("synthesis") or {}).get("analyses", {})
    if analyses:
        quality_el = ET.SubElement(root, "QUALITY_OF_EVIDENCE")
        for measure, analysis in analyses.items():
            grade = analysis.get("grade", {})
            if not grade:
                continue
            grade_el = ET.SubElement(quality_el, "GRADE_ASSESSMENT")
            grade_el.set("OUTCOME", measure)
            ET.SubElement(grade_el, "STARTING_LEVEL").text = grade.get("starting_level", "")
            ET.SubElement(grade_el, "FINAL_LEVEL").text = grade.get("final_level", "")
            for domain_name, value in grade.get("domains", {}).items():
                d_el = ET.SubElement(grade_el, "DOMAIN")
                d_el.set("NAME", domain_name)
                d_el.text = str(value)

    # Format nicely
    xml_str = ET.tostring(root, encoding="unicode", xml_declaration=False)
    try:
        pretty = minidom.parseString(xml_str).toprettyxml(indent="  ", encoding=None)
        # Remove extra xml declaration from minidom
        lines = pretty.split("\n")
        if lines and lines[0].startswith("<?xml"):
            lines = lines[1:]
        xml_str = '<?xml version="1.0" encoding="UTF-8"?>\n' + "\n".join(lines)
    except Exception:
        xml_str = '<?xml version="1.0" encoding="UTF-8"?>\n' + xml_str

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(xml_str)

    logger.info("RevMan XML exported: %s (%d studies)", output_path, len(papers))
    return output_path


# ══════════════════════════════════════════════════════════════════════════════
# 4. Summary Table (Markdown)
# ══════════════════════════════════════════════════════════════════════════════

def export_summary_table(pipeline_data, output_path):
    """Generate a Table of Characteristics in Markdown format.

    Standard systematic review table: Author/Year, Design, Population,
    Intervention, Comparator, Outcomes, Key Findings.

    Args:
        pipeline_data: Full pipeline output dict.
        output_path: Path for the .md file.

    Returns:
        str: The output path.
    """
    papers = _get_included_papers(pipeline_data)
    extraction_map = _get_extraction_map(pipeline_data)
    quality_map = _get_quality_map(pipeline_data)

    lines = []
    lines.append("# Table of Characteristics of Included Studies")
    lines.append("")
    lines.append("| Author/Year | Design | Population | Intervention | Comparator | Outcomes | Key Findings |")
    lines.append("|-------------|--------|------------|--------------|------------|----------|--------------|")

    for p in papers:
        pmid = str(p.get("pmid", ""))
        ext = extraction_map.get(pmid, {})
        qual = quality_map.get(pmid, {})

        sc = ext.get("study_characteristics", {})
        iv = ext.get("intervention_exposure", ext.get("intervention", {}))
        oc_raw = ext.get("outcomes", [])
        oc = oc_raw[0] if isinstance(oc_raw, list) and oc_raw else (oc_raw if isinstance(oc_raw, dict) else {})

        # Author/Year
        authors = p.get("authors", [])
        first_author = authors[0] if authors else "Unknown"
        year = p.get("year", sc.get("year", ""))
        author_year = f"{first_author} ({year})" if year else first_author

        design = sc.get("design", sc.get("study_design", "NR"))
        population = sc.get("population", "NR")
        intervention = iv.get("name", iv.get("intervention_name", "NR"))
        comparator = iv.get("comparator", iv.get("control", "NR"))
        outcome = oc.get("name", oc.get("primary_outcome", oc.get("outcome", "NR")))

        # Key findings: effect size + quality
        es = oc.get("effect_size_value", oc.get("effect_size", ""))
        es_type = oc.get("effect_size_type", oc.get("measure_type", ""))
        ci_lo = oc.get("ci_lower", "")
        ci_hi = oc.get("ci_upper", "")
        rob = qual.get("overall_judgment", qual.get("overall", ""))

        findings_parts = []
        if es and es_type:
            findings_parts.append(f"{es_type}={es}")
        if ci_lo and ci_hi:
            findings_parts.append(f"95%CI [{ci_lo}, {ci_hi}]")
        if rob:
            findings_parts.append(f"RoB: {rob}")
        findings = "; ".join(findings_parts) if findings_parts else "NR"

        # Escape pipes in cell content
        cells = [author_year, design, population, intervention, comparator, outcome, findings]
        cells = [str(c).replace("|", "/") for c in cells]

        lines.append(f"| {' | '.join(cells)} |")

    lines.append("")
    lines.append(f"*NR = Not Reported. Generated by SAM on {time.strftime('%Y-%m-%d %H:%M')}.*")
    lines.append("")

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))

    logger.info("Summary table exported: %s (%d studies)", output_path, len(papers))
    return output_path
